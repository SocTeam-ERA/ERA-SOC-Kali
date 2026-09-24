#!/usr/bin/env python3
"""
ad_report.py
------------
The monthly Active Directory account review as an Excel workbook, for the delegated AD admin and the
periodic review with IT: which computers and users are active, stale or disabled, who holds privileged
rights, which Windows versions have no support left, and the open domain weaknesses.

It reads only what ad_inventory.py already keeps (through ad_views.py), so the statuses are the same ones
the alerts and GET /api/ad/* use:
    active     signed in within ad_inventory.STALE_DAYS, or seen on the network recently
    stale      enabled but unused for longer: candidate to disable
    disabled   already disabled: candidate to delete
    review     system, service, privileged or server accounts: never clean up without the domain admin

Output: data/reports/ad/AD_review_<date>.xlsx (data/ is backed up, not in git: it names staff).
soc-ad-report.timer runs it on the 1st of each month.

    python3 ad_report.py                 write this month's workbook
    python3 ad_report.py --out FILE      write it somewhere else
"""
from __future__ import annotations

import argparse
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

sys.path.insert(0, str(Path(__file__).resolve().parent))
import ad_views  # noqa: E402
import soc_core  # noqa: E402

OUT_DIR = soc_core.DATA_DIR / "reports" / "ad"
NETWORK_RECENT_DAYS = 30

# Accounts that must never be cleaned up casually, whatever their activity.
SYSTEM_USERS = {"administrator", "krbtgt", "guest", "defaultaccount", "wdagutilityaccount"}
SERVICE_PREFIXES = ("svc", "msol_", "sync", "adsync", "qbdataservice")
NOTES = {
    "krbtgt": "Never delete: the domain's own Kerberos account",
    "svc-soc-ldap": "SOC read-only account: do not delete",
}


def _d(iso: Optional[str]) -> Optional[date]:
    return datetime.fromisoformat(iso).date() if iso else None


def _on_leave() -> Dict[str, str]:
    """{account: comment} from the on_leave_accounts watchlist (the comment line above each entry)."""
    import watchlists
    out, note = {}, ""
    path = watchlists._spec("on_leave_accounts")["path"]
    try:
        lines = path.read_text().splitlines()
    except OSError:
        return out
    for line in lines:
        line = line.strip()
        if line.startswith("#"):
            note = line.lstrip("# ").strip()
        elif line:
            out[line.lower()] = note
            note = ""
    return out


def computer_rows(today: date) -> List[Dict[str, Any]]:
    rows = []
    for c in ad_views.computers(today=today) or []:
        kind = "domain controller" if "Domain Controllers" in (c.get("ou") or "") else \
            "server" if "Server" in str(c.get("os") or "") else "system" if c["name"] == "AZUREADSSOACC" else ""
        net = c.get("network") or {}
        seen = _d(net.get("last_seen"))
        status = c["status"]
        if kind:
            status = "review"
        elif status == "stale" and seen and (today - seen).days <= NETWORK_RECENT_DAYS:
            status = "active"                   # AD's date lags, but the network saw it this month
        sup = c.get("support") or {}
        rows.append({"Computer": c["name"], "Status": status, "Kind": kind or "workstation",
                     "Enabled": "yes" if c.get("enabled") else "no", "Last sign-in (AD)": _d(c.get("last_logon")),
                     "Seen on network": seen, "IP": net.get("ip") or "", "OU": c.get("ou") or "",
                     "OS": c.get("os") or "", "Build": c.get("build"),
                     "Windows support": sup.get("status", ""), "Support ends": sup.get("ends", ""),
                     "Created": _d(c.get("created")), "Action": "", "Notes": ""})
    return rows


def user_rows(today: date) -> List[Dict[str, Any]]:
    leave = _on_leave()
    rows = []
    for u in ad_views.users(today=today) or []:
        key = u["sam"].lower()
        kind = "system" if key in SYSTEM_USERS else "privileged" if u["privileged_groups"] else \
            "service" if key.startswith(SERVICE_PREFIXES) else "person"
        note = NOTES.get(key, "") or ("Never delete: Microsoft 365 directory sync" if key.startswith(("msol_", "sync", "adsync"))
                                      else "")
        if key in leave:
            note = f"On leave, do not delete. {leave[key]}".strip()
        rows.append({"Account": u["sam"], "Name": u.get("display") or "",
                     "Status": "review" if kind != "person" else u["status"], "Kind": kind,
                     "Enabled": "yes" if u.get("enabled") else "no", "Last sign-in (AD)": _d(u.get("last_logon")),
                     "Created": _d(u.get("created")), "OU": u.get("ou") or "",
                     "Privileged groups": ", ".join(u["privileged_groups"]), "Action": "", "Notes": note})
    return rows


def build(out: Path, today: Optional[date] = None) -> Dict[str, int]:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    today = today or date.today()
    summary = ad_views.summary(today=today)
    if summary is None:
        raise SystemExit("no Active Directory inventory yet (ad_inventory.py)")
    comps, users = computer_rows(today), user_rows(today)

    font = "Arial"
    head_font, body = Font(name=font, bold=True, color="FFFFFF", size=10), Font(name=font, size=10)
    head_fill, edit_fill = PatternFill("solid", start_color="1F3864"), PatternFill("solid", start_color="FFF2CC")
    colors = {"active": "C6EFCE", "stale": "FFC7CE", "disabled": "E7E6E6", "review": "FFEB9C"}

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    def table(sheet, rows: List[Dict[str, Any]], widths: Dict[str, int]):
        cols = list(rows[0]) if rows else list(widths)
        for j, name in enumerate(cols, 1):
            c = sheet.cell(1, j, name)
            c.font, c.fill, c.alignment = head_font, head_fill, Alignment(wrap_text=True, vertical="center")
            sheet.column_dimensions[get_column_letter(j)].width = widths.get(name, 14)
        for i, r in enumerate(rows, 2):
            for j, name in enumerate(cols, 1):
                c = sheet.cell(i, j, r[name])
                c.font = body
                if isinstance(r[name], date):
                    c.number_format = "yyyy-mm-dd"
                if name in ("Action", "Notes"):
                    c.fill = edit_fill
        last = len(rows) + 1
        sheet.freeze_panes = "B2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(last, 1)}"
        if "Action" in cols and rows:
            col = get_column_letter(cols.index("Action") + 1)
            dv = DataValidation(type="list", formula1='"Keep,Disable,Delete"', allow_blank=True)
            sheet.add_data_validation(dv)
            dv.add(f"{col}2:{col}{last}")
        if "Status" in cols and rows:
            col = get_column_letter(cols.index("Status") + 1)
            for status, color in colors.items():
                sheet.conditional_formatting.add(f"{col}2:{col}{last}", FormulaRule(
                    formula=[f'${col}2="{status}"'], fill=PatternFill("solid", start_color=color)))
        return cols

    wc = wb.create_sheet("Computers")
    ccols = table(wc, comps, {"Computer": 18, "Status": 10, "Kind": 16, "Enabled": 8, "Last sign-in (AD)": 13,
                              "Seen on network": 13, "IP": 14, "OU": 24, "OS": 28, "Build": 8, "Windows support": 13,
                              "Support ends": 12, "Created": 12, "Action": 11, "Notes": 40})
    wu = wb.create_sheet("Users")
    ucols = table(wu, users, {"Account": 20, "Name": 26, "Status": 10, "Kind": 11, "Enabled": 8,
                              "Last sign-in (AD)": 13, "Created": 12, "OU": 24, "Privileged groups": 34,
                              "Action": 11, "Notes": 50})
    admins = [{"Group": g, "Account": m["sam"], "Name": m.get("display") or "", "Status": m.get("status") or "",
               "Last sign-in (AD)": _d(m.get("last_logon"))} for g, ms in summary["privileged"].items() for m in ms]
    table(wb.create_sheet("Admins"), admins, {"Group": 26, "Account": 20, "Name": 26, "Status": 10,
                                              "Last sign-in (AD)": 13})
    risks = [{"Severity": r["severity"], "Weakness": r["title"], "Accounts": ", ".join(r["accounts"][:40])
              + (f" (+{len(r['accounts']) - 40})" if len(r["accounts"]) > 40 else ""), "What to do": r["description"]}
             for r in summary.get("risks", [])]
    wr = table(wb.create_sheet("Domain weaknesses"), risks, {"Severity": 10, "Weakness": 50, "Accounts": 50,
                                                             "What to do": 90})
    for row in wb["Domain weaknesses"].iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    # Summary: counts are COUNTIFS over the sheets, so they follow any status a reviewer edits
    ws["A1"] = f"Active Directory review: {summary.get('as_of', '')[:10]}"
    ws["A1"].font = Font(name=font, bold=True, size=14)
    ws["A2"] = (f"From the SOC's daily AD inventory (server {summary.get('server')}). 'Stale' = enabled but unused for "
                f"90+ days. Disable first, delete after 30 days if nobody asks. People on leave are marked in Notes. Staff "
                f"who only use Microsoft 365 (mail, Teams) never sign in to AD and look stale: deleting their AD account "
                f"deletes their mailbox, so check first.")
    ws["A2"].font = Font(name=font, italic=True, size=9)
    st_col_c = get_column_letter(ccols.index("Status") + 1)
    st_col_u = get_column_letter(ucols.index("Status") + 1)
    r = 4
    for title, sheet, col in (("Computers", "Computers", st_col_c), ("Users", "Users", st_col_u)):
        ws.cell(r, 1, title).font = Font(name=font, bold=True, size=11)
        r += 1
        start = r
        for status in ("active", "stale", "disabled", "review"):
            ws.cell(r, 1, status).font = body
            ws.cell(r, 2, f'=COUNTIF({sheet}!${col}:${col},"{status}")').font = body
            r += 1
        ws.cell(r, 1, "total").font = Font(name=font, bold=True, size=10)
        ws.cell(r, 2, f"=SUM(B{start}:B{r - 1})").font = Font(name=font, bold=True, size=10)
        r += 2
    ws.cell(r, 1, "Windows without support (active PCs)").font = body
    ws.cell(r, 2, len(summary["windows_support"]["unsupported"])).font = body
    ws.cell(r + 1, 1, "Losing support within 45 days").font = body
    ws.cell(r + 1, 2, len(summary["windows_support"]["ending_soon"])).font = body
    ws.cell(r + 2, 1, "Windows machines outside the domain").font = body
    ws.cell(r + 2, 2, len(summary["not_in_domain"])).font = body
    ws.cell(r + 3, 1, "Open domain weaknesses").font = body
    ws.cell(r + 3, 2, len(summary.get("risks", []))).font = body
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 40, 10
    wb.calculation.fullCalcOnLoad = True

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return {"computers": len(comps), "users": len(users), "admins": len(admins), "risks": len(risks)}


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--out", type=Path, help="output file (default data/reports/ad/AD_review_<date>.xlsx)")
    args = ap.parse_args()
    target = args.out or OUT_DIR / f"AD_review_{date.today().isoformat()}.xlsx"
    counts = build(target)
    print(f"[*] ad_report: {target} ({counts})")

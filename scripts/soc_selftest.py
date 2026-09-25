#!/usr/bin/env python3
"""
soc_selftest.py -- does the whole detection pipeline still work? Answer in about a minute.

Three parts:

  1. DETECTORS (isolated). Each detector is fed a small synthetic input it should react to
     (a Suricata scan alert, a Zeek password-guessing notice, five failed SSH logins, an
     osquery new-user row, a wave of ARP new devices, an nmap XML with a new port, an EICAR
     test file, a phishing email, tshark rows with a port scan...) and the resulting alert is
     checked. It all runs in a throwaway data directory with notifications and the backend
     forwarder switched off, so nothing reaches the dashboard or anyone's phone.
     Then the platform layers on top: MITRE tags, entities, batches, threat intel,
     correlation into an incident, and suppression; and the analyst-facing layers: alert
     aging, search, incident lifecycle, suppressions created from the dashboard, playbooks,
     the backup script, and an isolated API server exercised with read and write keys.

  2. LIVE (read-only). Services and timers active, the API answering with a real key, the
     data sources healthy, the backup and threat-intel feeds fresh, the config files valid,
     and soc_doctor passing. Nothing is written.

  3. CANARY (only with --canary). Sends one harmless "normal" alert through the real pipeline,
     checks that it shows up in the live feed and through the API, then resolves it.

    python3 soc_selftest.py            parts 1 and 2
    python3 soc_selftest.py --canary   parts 1, 2 and 3
    python3 soc_selftest.py --isolated | --live      only one part
    python3 soc_selftest.py --live --alert-on-fail   (what the daily cron runs) raise ONE alert
                                                     when checks fail; a repeat of the same failures
                                                     is not re-alerted, and a recovery is recorded

Exit code 0 = everything passed. The summary is saved to data/selftest_last.json:
{at, passed, total, failed, failures: [{name, detail}]}, which soc_doctor reads.
It only uses harmless test inputs (the EICAR string is the industry-standard antivirus
test file) and never touches the network beyond this machine's own API.
"""
from __future__ import annotations

import json
import re
import os
import subprocess
import sys
import tempfile
import time
import urllib.parse
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

SCRIPTS = Path(__file__).resolve().parent
SUITE = SCRIPTS.parent
KALI = SUITE / "kali"
DATA = Path(os.environ.get("SOC_DATA_DIR", SUITE / "data"))
API = os.environ.get("SOC_API_URL", "http://127.0.0.1:8080")

results: list[tuple[str, bool, str]] = []
current_group = ""


def check(name: str, ok: bool, detail: str = "") -> None:
    results.append((f"{current_group}: {name}" if current_group else name, bool(ok), detail))


def group(name: str) -> None:
    global current_group
    current_group = name


# --------------------------------------------------------------------------- #
#  Part 1 -- detectors and platform layers, in a throwaway data dir
# --------------------------------------------------------------------------- #

def inner() -> int:
    """Runs with SOC_DATA_DIR pointing at a temp dir (set by run_isolated)."""
    sys.path[:0] = [str(KALI), str(SCRIPTS)]
    import soc_core
    tmp = Path(os.environ["SOC_DATA_DIR"])

    def feed() -> list:
        try:
            return json.loads(soc_core.ALERTS_SNAPSHOT.read_text())
        except (OSError, ValueError):
            return []

    def new_alerts(before: int) -> list:
        return feed()[: max(len(feed()) - before, 0)]

    def find(alerts: list, text: str, **kw):
        for a in alerts:
            if text in a["title"] and all(a.get(k) == v for k, v in kw.items()):
                return a
        return None

    def tags(a) -> set:
        return {t["technique"] for t in (a or {}).get("details", {}).get("mitre", [])}

    # ---- Suricata ------------------------------------------------------------
    group("suricata")
    import suricata_to_alerts as S
    n = len(feed())
    ok = S.handle_event({"event_type": "alert", "src_ip": "203.0.113.9", "dest_ip": "10.201.5.5", "dest_port": 80,
                         "proto": "TCP", "alert": {"severity": 1, "signature": "ET SCAN Possible Nmap User-Agent Observed",
                                                   "category": "Web Application Attack", "signature_id": 2024364, "gid": 1}})
    a = find(new_alerts(n), "ET SCAN")
    check("a scan signature becomes an alert", ok and a is not None, str(ok))
    check("tagged MITRE T1046 and has entities", "T1046" in tags(a) and bool((a or {}).get("details", {}).get("entities")))
    check("non-alert events (flows) are ignored", S.handle_event({"event_type": "flow", "src_ip": "10.0.0.1"}) is False)

    # ---- Zeek ----------------------------------------------------------------
    group("zeek")
    import zeek_to_alerts as Z
    n = len(feed())
    ok = Z.handle_event({"note": "SSH::Password_Guessing", "msg": "203.0.113.9 appears to be guessing SSH passwords",
                         "src": "203.0.113.9", "dst": "10.201.5.5"})
    a = find(new_alerts(n), "SSH::Password_Guessing")
    check("a password-guessing notice becomes an alert", ok and a is not None)
    check("tagged MITRE T1110.001", "T1110.001" in tags(a))
    check("known-noise notices are dropped (CaptureLoss)", Z.handle_event({"note": "CaptureLoss::Too_Much_Loss", "msg": "x"}) is False)
    zt = tmp / "notice_tsv.log"
    zt.write_text("#separator \\x09\n#fields\tts\tnote\tmsg\tsrc\tdst\tp\n#types\ttime\tenum\tstring\taddr\taddr\tport\n"
                  "1.0\tSSH::Password_Guessing\ttsv brute force from 203.0.113.60\t203.0.113.60\t10.201.5.5\t22\n"
                  "2.0\tCaptureLoss::Too_Much_Loss\tnoise\t-\t-\t-\n")
    n = len(feed())
    forwarded = Z.process_file(zt, follow=False)
    a = find(new_alerts(n), "tsv brute force")
    check("a notice in Zeek's classic TSV format becomes an alert, its port typed as a number (regression: this forwarder "
          "json.loads()-ed TSV lines and forwarded nothing for two days after Zeek's format flipped)",
          forwarded == 1 and a is not None and a["severity"] == "critical" and a["details"]["p"] == 22)
    zj = tmp / "notice_json.log"
    zj.write_text(json.dumps({"ts": 1.0, "note": "SSH::Password_Guessing", "msg": "json brute force from 203.0.113.61",
                              "src": "203.0.113.61", "dst": "10.201.5.5", "p": 22}) + "\n")
    n = len(feed())
    check("...and the JSON format still works", Z.process_file(zj, follow=False) == 1 and find(new_alerts(n), "json brute force"))

    group("chkrootkit")
    import chkrootkit_to_alerts as CK
    check("this appliance's own nmap is not an unrecognized packet sniffer (it fired every night: the cron.daily job "
          "runs mid-scan)", CK.unexplained_ifpromisc(["eth2: PACKET SNIFFER(/opt/zeek/bin/zeek[1604], /usr/lib/nmap/nmap[39680])"]) == [])
    check("...but a sniffer nobody expects still is one",
          CK.unexplained_ifpromisc(["eth2: PACKET SNIFFER(/tmp/x/evil-sniffer[1])"]) == ["eth2: unrecognized sniffer(s) evil-sniffer"])

    # ---- osquery -------------------------------------------------------------
    group("osquery")
    import osquery_to_alerts as O
    seen: dict = {}
    n = len(feed())
    base = {"name": "local_users", "action": "added", "counter": 1, "columns": {"username": "alice", "uid": "1000"}}
    check("the first batch only records a baseline (no alert)", O.handle_row(base, seen) is False)
    ok = O.handle_row({"name": "local_users", "action": "added", "counter": 2,
                       "columns": {"username": "backdoor", "uid": "0"}}, seen)
    a = find(new_alerts(n), "backdoor")
    check("a NEW local user after the baseline is a critical alert", ok and a is not None and a["severity"] == "critical")

    # ---- login monitor -------------------------------------------------------
    group("login_monitor")
    import ipaddress
    import login_monitor as L
    trk = L.BruteForceTracker(5, 60)
    known = [ipaddress.ip_network("10.0.0.0/8")]
    parsed = L.parse_line("Sep 21 10:00:01 kali2 sshd[123]: Failed password for root from 203.0.113.50 port 5555 ssh2")
    check("an sshd failure line is parsed", parsed is not None and parsed[2] == "203.0.113.50", str(parsed))
    n = len(feed())
    for _ in range(5):
        L.handle_event("failed", "root", "203.0.113.50", None, trk, known)
    a = find(new_alerts(n), "Brute-force")
    check("5 failures from one IP raise a brute-force alert", a is not None)
    n = len(feed())
    L.handle_event("accepted", "root", "203.0.113.50", None, trk, known)
    check("a login right after the failures is flagged as possible compromise",
          find(new_alerts(n), "Successful login after") is not None)
    n = len(feed())
    L.handle_event("accepted", "carol", "198.51.100.7", None, trk, known)
    check("a login from an IP outside the trusted list is flagged", find(new_alerts(n), "Login from NEW IP") is not None)
    n = len(feed())
    L.handle_event("accepted", "carol", "198.51.100.7", None, trk, known)
    check("...and only once per IP and user (cooldown)", find(new_alerts(n), "Login from NEW IP") is None)
    kf = tmp / "known_ips_test.txt"
    kf.write_text("10.0.0.0/8\n")
    kn = L.load_known_ips(str(kf))
    check("the trusted-IP list is loaded from its file", L.is_known_ip("10.1.2.3", kn) and not L.is_known_ip("198.51.100.77", kn))
    with kf.open("a") as fh:
        fh.write("198.51.100.77\n")
    check("an address added to that file later is trusted without a restart (regression: it was read once at startup, "
          "so entries added through the watchlist API never took effect)", L.is_known_ip("198.51.100.77", kn))

    # ---- AIDE ----------------------------------------------------------------
    group("aide")
    import aide_to_alerts as A
    report = ("Added entries:\n---------------------------------------------------\n\n"
              "f++++++++++++++++++: /etc/cron.d/evil\n\n"
              "Removed entries:\n---------------------------------------------------\n\n"
              "f------------------: /usr/local/bin/old\n\n"
              "Changed entries:\n---------------------------------------------------\n\n"
              "f =.... mc..H.. .  : /etc/passwd\n")
    entries = dict((p, k) for k, p in A.parse_report(report))
    check("added, removed and changed entries are parsed",
          entries.get("/etc/cron.d/evil") == "added" and entries.get("/usr/local/bin/old") == "removed"
          and entries.get("/etc/passwd") == "changed", str(entries))
    check("changes to /etc/passwd are critical", A.severity_for("/etc/passwd") == "critical")
    from soc_core import Alert, emit_alert
    n = len(feed())
    emit_alert(Alert(type="intrusion", severity="critical", title="File integrity: /etc/cron.d/evil appeared",
                     detector="aide", details={"path": "/etc/cron.d/evil", "change": "added"}), echo=False)
    check("a new cron file is tagged MITRE T1053.003", "T1053.003" in tags(feed()[0]))

    # ---- traffic monitor -----------------------------------------------------
    group("traffic")
    import traffic_to_alerts as T
    own = sorted(T._own_ips())
    rows = []
    for p in range(1000, 1040):                                   # 40 distinct ports from one external host
        rows.append(f"1\t203.0.113.77\t10.201.5.5\t{p}\t\ttcp\t\t\t1\t0")
    for _ in range(3):                                            # cleartext Telnet, real session (not bare SYN)
        rows.append("2\t10.201.5.20\t10.201.0.13\t23\t\ttcp\t\t\t0\t1")
    rows.append("3\t10.201.5.30\t185.220.101.4\t443\t\ttcp\t\t\t0\t1")   # known-bad IP
    rows.append("4\t10.201.5.31\t8.8.4.4\t\t53\tudp\ttotally-legit-updates.xyz\t\t\t")   # bad domain
    if own:
        rows += [f"5\t{own[0]}\t10.201.5.5\t{p}\t\ttcp\t\t\t1\t0" for p in range(2000, 2020)]   # 20 ports from THIS box
    n = len(feed())
    T.parse(iter([r + "\n" for r in rows]), {"185.220.101.4"}, {"totally-legit-updates.xyz"}, None, set())
    new = new_alerts(n)
    check("a 40-port sweep from another host is a port scan", find(new, "Port scan on the wire: 203.0.113.77") is not None)
    check("Telnet in use is flagged as a cleartext protocol", find(new, "Cleartext protocol Telnet") is not None)
    check("contact with a known-bad IP is flagged", find(new, "known-bad IP 185.220.101.4") is not None)
    check("a query for a bad domain is flagged", find(new, "totally-legit-updates.xyz") is not None)
    if own:
        check("this appliance's own 20-port traffic is NOT a scan (own-IP threshold)",
              find(new, f"Port scan on the wire: {own[0]}") is None)

    rows_dns = [f"6\t10.69.0.14\t{own[0] if own else '10.201.5.5'}\t\t{p}\tudp\t\t\t\t"
                for p in range(40000, 40045)]   # 45 "destination ports" all in the ephemeral range
    n = len(feed())
    T.parse(iter([r + "\n" for r in rows_dns]), set(), set(), None, set())
    check("a DNS server answering many of our own ephemeral ports is NOT a scan",
          find(new_alerts(n), "Port scan on the wire: 10.69.0.14") is None)

    # ---- ARP discovery -------------------------------------------------------
    group("arp")
    import arp_to_alerts as R
    def tsv(rows_):
        p = tmp / "arp.tsv"; p.write_text("\n".join("\t".join(r) for r in rows_) + "\n"); return p
    base_rows = [("10.69.0.0/16", "eth0", f"10.69.9.{i}", f"3c:bb:cc:00:00:{i:02x}", "Dell Inc.") for i in range(1, 4)]
    R.run(tsv(base_rows), tmp / "mac_state.json")
    n = len(feed())
    R.run(tsv(base_rows + [("10.69.0.0/16", "eth0", "10.69.8.1", "3c:bb:cc:00:01:01", "Intel Corporate")]), tmp / "mac_state.json")
    a = find(new_alerts(n), "New device on VLAN")
    check("one new device raises a new-device alert", a is not None and a["severity"] == "medium")
    check("tagged MITRE T1200 and carries its MAC as an entity", "T1200" in tags(a)
          and any(e["type"] == "mac" for e in (a or {}).get("details", {}).get("entities", [])))
    n = len(feed())
    wave = base_rows + [("10.69.0.0/16", "eth0", f"10.69.7.{i}", f"3c:bb:cc:00:02:{i:02x}", "HP") for i in range(1, 9)]
    R.run(tsv(wave), tmp / "mac_state.json")
    wave_alerts = new_alerts(n)
    bulk = find(wave_alerts, "new device(s) on VLAN")
    check("a wave of 8 new devices becomes ONE grouped alert", bulk is not None and len(wave_alerts) == 1 and bulk["details"]["count"] == 8)

    # ---- (added) asset enrichment: Zeek TSV parsing, DHCP hostnames, software fingerprints ----
    group("asset enrichment")
    import zeek_tsv
    r = zeek_tsv.ZeekTSVReader()
    check("a data row before any header is dropped", r.feed("2026\tfoo\tbar") is None and not r.ready)
    check("a #fields header is recognized and not returned as data", r.feed("#fields\tts\thost\tname") is None and r.ready)
    check("a data row after the header is parsed into a dict",
          r.feed("123.0\t10.0.0.5\tGoAhead-Webs") == {"ts": "123.0", "host": "10.0.0.5", "name": "GoAhead-Webs"})
    check("Zeek's unset marker '-' becomes None", r.feed("123.0\t-\tfoo") == {"ts": "123.0", "host": None, "name": "foo"})
    check("a #types/#open/#close comment line is ignored", r.feed("#types\ttime\taddr\tstring") is None)
    check("a malformed row (wrong column count) is dropped, not misaligned", r.feed("123.0\tonly-two") is None)
    rj = zeek_tsv.ZeekTSVReader()
    check("a JSON-format line parses with no header at all", rj.feed('{"ts":1.5,"note":"X::Y","p":22}') == {"ts": 1.5, "note": "X::Y", "p": 22})
    check("a malformed JSON line is dropped", rj.feed('{"ts":') is None)
    rt = zeek_tsv.ZeekTSVReader()
    rt.feed("#fields\tts\tid.resp_p\ttags\tok\tn")
    rt.feed("#types\ttime\tport\tset[string]\tbool\tcount")
    check("typed columns come out like Zeek's JSON form (time float, port int, set list, bool, count int)",
          rt.feed("1.5\t443\ta,b\tT\t7") == {"ts": 1.5, "id.resp_p": 443, "tags": ["a", "b"], "ok": True, "n": 7})
    check("an unset field is None and an (empty) set is []",
          rt.feed("1.5\t-\t(empty)\tF\t-") == {"ts": 1.5, "id.resp_p": None, "tags": [], "ok": False, "n": None})
    import gzip
    gz_path = tmp / "hdr_test.log.gz"
    with gzip.open(gz_path, "wt") as gz:
        gz.write("#separator \\x09\n#fields\tts\thost\n#types\ttime\taddr\n1.0\t10.0.0.9\n")
    check("read_header_lines reads the #fields and #types lines of a .gz archive too",
          zeek_tsv.read_header_lines(gz_path) == ["#fields\tts\thost", "#types\ttime\taddr"])

    hdr_file = tmp / "hdr_test.log"
    hdr_file.write_text("#separator \\x09\n#fields\tts\thost\n#types\ttime\taddr\n1.0\t10.0.0.9\n")
    check("read_current_header finds the #fields line of a real file",
          zeek_tsv.read_current_header(hdr_file) == "#fields\tts\thost")
    no_hdr = tmp / "no_header.log"; no_hdr.write_text("just data, no header\n")
    check("read_current_header returns None when there is no header", zeek_tsv.read_current_header(no_hdr) is None)

    ipmap = soc_core.build_ip_to_mac_map()
    check("build_ip_to_mac_map reflects the asset inventory", ipmap.get("10.69.9.1") == "3c:bb:cc:00:00:01")

    changed = soc_core.record_asset_software([{"mac": "3c:bb:cc:00:00:01", "software_type": "HTTP::SERVER",
                                               "name": "GoAhead-Webs", "version": "GoAhead-Webs"}])
    check("record_asset_software enriches an existing asset", changed == 1
          and soc_core.load_assets()["3c:bb:cc:00:00:01"]["software"]["HTTP::SERVER"]["name"] == "GoAhead-Webs")
    changed = soc_core.record_asset_software([{"mac": "3c:bb:cc:00:00:01", "software_type": "HTTP::SERVER",
                                               "name": "GoAhead-Webs", "version": "GoAhead-Webs"}])
    check("...and an identical repeat sighting is a no-op", changed == 0)
    changed = soc_core.record_asset_software([{"mac": "00:00:00:00:00:99", "software_type": "HTTP::SERVER",
                                               "name": "x", "version": None}])
    check("record_asset_software never creates a new asset record",
          changed == 0 and "00:00:00:00:00:99" not in soc_core.load_assets())

    import dhcp_to_assets as DA
    dhcp_log = tmp / "dhcp_test.log"
    dhcp_log.write_text(
        "#separator \\x09\n#fields\tts\tmac\thost_name\tclient_fqdn\n#types\ttime\tstring\tstring\tstring\n"
        "1.0\t3c:bb:cc:00:00:02\tMYPRINTER\t-\n"
        "2.0\taa:aa:aa:aa:aa:aa\tGHOST\t-\n")  # not a known asset -- must be skipped
    n = DA.process_file(dhcp_log, follow=False)
    check("dhcp_to_assets parses the real Zeek TSV format and enriches a known asset (regression: "
          "this used to json.loads() a TSV line and silently enrich nothing, ever)",
          n == 1 and soc_core.load_assets()["3c:bb:cc:00:00:02"].get("dhcp_hostname") == "MYPRINTER")

    import software_to_assets as SA
    sw_log = tmp / "software_test.log"
    sw_log.write_text(
        "#separator \\x09\n#fields\tts\thost\thost_p\tsoftware_type\tname\tunparsed_version\n"
        "#types\ttime\taddr\tport\tenum\tstring\tstring\n"
        "1.0\t10.69.9.3\t22\tSSH::SERVER\tOpenSSH\tOpenSSH_9.1\n"
        "2.0\t203.0.113.9\t80\tHTTP::SERVER\tnginx\tnginx/1.18\n")  # not a known asset -- must be skipped
    n = SA.process_file(sw_log, follow=False)
    check("software_to_assets resolves IP to MAC via the asset inventory and enriches a known asset",
          n == 1 and soc_core.load_assets()["3c:bb:cc:00:00:03"]["software"]["SSH::SERVER"]["version"] == "OpenSSH_9.1")

    dj = tmp / "dhcp_json.log"
    dj.write_text(json.dumps({"ts": 1.0, "mac": "3c:bb:cc:00:00:01", "host_name": "JSONHOST"}) + "\n")
    check("dhcp_to_assets also reads the JSON log format", DA.process_file(dj, follow=False) == 1
          and soc_core.load_assets()["3c:bb:cc:00:00:01"].get("dhcp_hostname") == "JSONHOST")
    sj = tmp / "software_json.log"
    sj.write_text(json.dumps({"ts": 1.0, "host": "10.69.9.1", "software_type": "HTTP::SERVER", "name": "lighttpd",
                              "unparsed_version": "lighttpd/1.4"}) + "\n")
    check("software_to_assets also reads the JSON log format", SA.process_file(sj, follow=False) == 1
          and soc_core.load_assets()["3c:bb:cc:00:00:01"]["software"]["HTTP::SERVER"]["name"] == "lighttpd")

    late_notice, late_software = tmp / "late_notice.log", tmp / "late_software.log"
    followers = [subprocess.Popen([sys.executable, str(KALI / "zeek_to_alerts.py"), "--follow", "--log", str(late_notice)],
                                  stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL),
                 subprocess.Popen([sys.executable, str(KALI / "software_to_assets.py"), "--follow", "--log", str(late_software)],
                                  stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)]
    try:
        time.sleep(1.5)
        check("--follow waits for a log that does not exist yet instead of exiting (regression: the forwarder exited 1 and "
              "systemd restarted it in a loop whenever a quiet hour left no notice.log)",
              all(f.poll() is None for f in followers))
        late_notice.write_text("#separator \\x09\n#fields\tts\tnote\tmsg\tsrc\tdst\n#types\ttime\tenum\tstring\taddr\taddr\n"
                               "1.0\tSSH::Password_Guessing\tlate brute force from 203.0.113.62\t203.0.113.62\t10.201.5.5\n")
        late_software.write_text("#separator \\x09\n#fields\tts\thost\thost_p\tsoftware_type\tname\tunparsed_version\n"
                                 "#types\ttime\taddr\tport\tenum\tstring\tstring\n"
                                 "1.0\t10.69.9.2\t80\tHTTP::SERVER\tlighttpd\tlighttpd/1.4.71\n")
        deadline, got_alert, got_asset = time.time() + 8, False, False
        while time.time() < deadline and not (got_alert and got_asset):
            got_alert = got_alert or find(feed(), "late brute force") is not None
            got_asset = got_asset or "HTTP::SERVER" in soc_core.load_assets().get("3c:bb:cc:00:00:02", {}).get("software", {})
            time.sleep(0.3)
        check("a log that appears later is read from its first line, header included (not from its end)", got_alert and got_asset,
              f"notice alert: {got_alert}, software enrichment: {got_asset}")
        import reader_health
        published = {n_: reader_health.status(n_) for n_ in ("zeek_to_alerts", "software_to_assets")}
        check("a running follower publishes its parse health where source_health can read it (written on start, then at most "
              "once a minute)", all(v_ is not None and v_["streak"] == 0 for v_ in published.values()), str(published))
    finally:
        for f in followers:
            f.terminate()
        for f in followers:
            try:
                f.wait(timeout=5)
            except subprocess.TimeoutExpired:
                f.kill()



    # ---- nmap ----------------------------------------------------------------
    group("nmap")
    import nmap_to_alerts as N
    def xml(ports):
        body = "".join(f'<port protocol="tcp" portid="{p}"><state state="open"/><service name="{s}" method="probed" conf="10"/></port>'
                       for p, s in ports)
        return ('<?xml version="1.0"?><nmaprun><host><status state="up"/><address addr="10.1.1.5" addrtype="ipv4"/>'
                f'<hostnames/><ports>{body}</ports></host></nmaprun>')
    st = tmp / "port_state.json"
    (tmp / "scan1.xml").write_text(xml([(22, "ssh")])); N.run(tmp / "scan1.xml", {22}, st)
    n = len(feed())
    (tmp / "scan2.xml").write_text(xml([(22, "ssh"), (8080, "http")])); N.run(tmp / "scan2.xml", {22}, st)
    a = find(new_alerts(n), "NEW open port 8080/tcp")
    check("a port that opened between two scans raises an alert (normal: a new ordinary port is information)", a is not None and a["severity"] == "normal")
    n = len(feed()); N.run(tmp / "scan2.xml", {22}, st)
    check("the same scan again raises nothing (change detection)", len(new_alerts(n)) == 0)
    n = len(feed())
    (tmp / "scan3.xml").write_text(xml([(22, "ssh"), (8080, "http"), (3389, "ms-wbt-server")])); N.run(tmp / "scan3.xml", {22}, st)
    a = find(new_alerts(n), "NEW open port 3389/tcp")
    check("a new RDP port on a known device is medium and tagged T1021.001", a is not None and a["severity"] == "medium" and "T1021.001" in tags(a))

    # ---- malware / phishing --------------------------------------------------
    group("malware")
    import malware_detector as M
    eicar = tmp / "eicar.txt"
    eicar.write_text(r"X5O!P%@AP[4\PZX54(P^)7CC)7}$EICAR-STANDARD-ANTIVIRUS-TEST-FILE!$H+H*")
    ioc = tmp / "iocs.txt"; ioc.write_text("44d88612fea8a8f36de82e1278abb02f\n")
    n = len(feed()); M.scan_file(eicar, M.load_iocs(ioc))
    a = find(new_alerts(n), "Suspicious file")
    check("the EICAR test file is caught by its hash and is critical", a is not None and a["severity"] == "critical")
    group("phishing")
    import phishing_detector as P
    eml = tmp / "sample.eml"; eml.write_text(P.DEMO_EML)
    n = len(feed()); P.analyse_file(eml)
    check("the bundled phishing sample raises a phishing alert", find(new_alerts(n), "Phishing indicators") is not None)

    # ---- platform layers -----------------------------------------------------
    group("platform")
    import correlate
    n = len(feed())
    emit_alert(Alert(type="port_scan", severity="medium", title="Port scan on the wire: 192.0.2.44 probed 120 ports",
                     detector="traffic_capture", source_ip="192.0.2.44"), echo=False)
    emit_alert(Alert(type="vuln", severity="critical", title="ET EXPLOIT D-Link command injection (192.0.2.44 -> 10.201.5.5)",
                     detector="suricata", source_ip="192.0.2.44"), echo=False)
    incs = [i for i in correlate.list_incidents() if i["rule_id"] == "scan-then-exploit"]
    check("a scan followed by an exploit from the same host opens a critical incident",
          len(incs) == 1 and incs[0]["severity"] == "critical")
    check("the incident raised its own alert (detector correlation)", find(new_alerts(n), "Incident #") is not None)
    a = emit_alert(Alert(type="malware", severity="medium", title="Contact", detector="traffic_capture",
                         source_ip="185.220.101.4"), echo=False)
    check("a source in the local IOC list is enriched with threat intel", bool(a["details"].get("threat_intel")))
    a1 = emit_alert(Alert(type="intrusion", severity="normal", title="batch probe", detector="selftest"), echo=False)
    a2 = emit_alert(Alert(type="intrusion", severity="normal", title="batch probe", detector="selftest"), echo=False)
    check("alerts of one kind close together share a batch id", a1["details"]["batch_id"] == a2["details"]["batch_id"])
    supp = tmp / "sup.json"
    supp.write_text(json.dumps({"rules": [{"id": "selftest-rule", "reason": "self-test rule",
                                            "match": {"detector": "selftest_suppressed", "title_contains": "hide me"}}]}))
    import suppressions
    suppressions.CONFIG_FILE = supp
    r1 = emit_alert(Alert(type="intrusion", severity="medium", title="please hide me", detector="selftest_suppressed"), echo=False)
    r2 = emit_alert(Alert(type="intrusion", severity="medium", title="but not this", detector="selftest_suppressed"), echo=False)
    check("a matching suppression rule hides an alert, others pass", r1.get("suppressed_by") == "selftest-rule" and "suppressed_by" not in r2)

    # ---- (added) alert aging --------------------------------------------------
    group("aging")
    import alert_aging
    from datetime import timedelta

    def _ago(days):
        return (datetime.now(timezone.utc) - timedelta(days=days)).isoformat()
    old_port = emit_alert(Alert(type="port_scan", severity="medium", title="NEW open port 80/tcp (http) on 10.0.0.1",
                                detector="kali_scan", timestamp=_ago(6)), echo=False)
    old_crit = emit_alert(Alert(type="port_scan", severity="critical", title="NEW open port 3389/tcp on 10.0.0.2",
                                detector="kali_scan", timestamp=_ago(9)), echo=False)
    old_vuln = emit_alert(Alert(type="vuln", severity="medium", title="Weak TLS on 10.0.0.5:443 (ssl-cert): certificate EXPIRED",
                                detector="kali_scan", timestamp=_ago(6)), echo=False)
    new_port = emit_alert(Alert(type="port_scan", severity="medium", title="NEW open port 81/tcp on 10.0.0.3",
                                detector="kali_scan", timestamp=_ago(1)), echo=False)
    alert_aging.run()
    st = {a["id"]: a["status"] for a in feed()}
    check("a 6-day-old informational port alert is closed automatically", st.get(old_port["id"]) == "resolved")
    check("an old CRITICAL alert is left open", st.get(old_crit["id"]) == "open")
    check("an old vulnerability finding is left open", st.get(old_vuln["id"]) == "open")
    check("a recent port alert is left open", st.get(new_port["id"]) == "open")
    soc_core.set_alert_status(old_port["id"], "open", note="[jdoe] still investigating", actor="platform-backend")
    alert_aging.run()
    check("an alert a person reopened (e.g. from the dashboard) is not closed again by the aging timer",
          {a["id"]: a["status"] for a in feed()}.get(old_port["id"]) == "open")
    sm = soc_core.summarize()
    check("the summary also counts only what still needs a person (open + acknowledged) under 'active'",
          sm["active"]["total"] == sm["by_status"]["open"] + sm["by_status"]["acknowledged"]
          and sum(sm["active"]["by_severity"].values()) == sm["active"]["total"])

    # ---- (added) search -------------------------------------------------------
    group("search")
    import log_search
    res = log_search.search("alerts", "batch probe", "1h", 10, 5)
    check("a text search finds alerts in the history", res["count"] >= 1, str(res["count"]))
    res = log_search.search("alerts", "detector:kali_scan ip:10.0.0.0/8", "1h", 50, 5)
    check("a field + network query finds the port alerts", res["count"] >= 1, str(res["count"]))

    def _refused(fn):
        try:
            fn()
        except ValueError:
            return True
        return False
    check("a path-traversal source is refused", _refused(lambda: log_search.search("zeek:../../etc/passwd", "x")))
    check("an empty query is refused", _refused(lambda: log_search.search("alerts", "")))

    import gzip as _gzip
    zroot = tmp / "zeek"
    (zroot / "current").mkdir(parents=True)
    now_ = time.time()
    (zroot / "current" / "dns.log").write_text(
        "#separator \\x09\n#fields\tts\tid.orig_h\tid.resp_p\tquery\n#types\ttime\taddr\tport\tstring\n"
        f"{now_ - 60:.6f}\t10.1.1.5\t53\ttsv-example.test\n{now_ - 30:.6f}\t10.1.1.6\t53\tother.test\n#close\t2026-09-23\n")
    day = time.strftime("%Y-%m-%d")
    (zroot / day).mkdir()
    with _gzip.open(zroot / day / "notice.00:00:00-23:59:59.log.gz", "wt") as gz:
        gz.write(json.dumps({"ts": now_ - 120, "note": "X::FromArchive", "msg": "json archive row"}) + "\n")
    real_zeek_dir, log_search.ZEEK_DIR = log_search.ZEEK_DIR, zroot
    try:
        res = log_search.search("zeek:dns", "query~tsv-example dport:53", "1h", 10, 5)
        check("a Zeek log in TSV format is searchable, with typed fields (regression: this returned nothing for any Zeek data "
              "newer than the format flip)", res["count"] == 1 and res["results"][0]["record"]["id.orig_h"] == "10.1.1.5"
              and res["results"][0]["record"]["id.resp_p"] == 53, str(res["count"]))
        res = log_search.search("zeek:notice", "FromArchive", "1h", 10, 5)
        check("a JSON archive of a log that has no live file is searchable", res["count"] == 1, str(res["count"]))
        check("the source list includes logs that exist only as archives", "zeek:notice" in log_search.available_sources()["sources"])
        check("a made-up Zeek log name is still refused", _refused(lambda: log_search.search("zeek:nope", "x")))
    finally:
        log_search.ZEEK_DIR = real_zeek_dir

    # ---- (added) incidents ----------------------------------------------------
    group("incidents")
    incs = correlate.list_incidents()
    inc_id = incs[0]["id"] if incs else None
    check("an incident exists to work with", inc_id is not None)
    if inc_id:
        check("closing an incident without a classification is refused",
              _refused(lambda: correlate.update_incident(inc_id, "selftest", status="closed")))
        upd = correlate.update_incident(inc_id, "selftest", status="closed", classification="benign", comment="self-test")
        check("it closes with a classification and keeps the comment", upd["status"] == "closed" and len(upd["comments"]) == 1)
        check("the change is attributed to who made it", upd["comments"][0]["by"] == "selftest")

    # the rules that use what the Kali sees without a mirror port (config/correlation_rules.json)
    import alert_match
    rules_, _, _ = correlate.load_rules()
    byid = {r_["id"]: r_ for r_ in rules_}
    check("the four rules for new devices, the AD inventory and exposure load, with 7-day windows allowed",
          {"unmanaged-windows-device", "new-device-exposes-services", "exposed-unpatched-host",
           "unmanaged-host-scanning"} <= set(byid))
    nd = byid.get("new-device-exposes-services", {}).get("steps", [{}])[0]
    mt = lambda title, det: alert_match.matches(nd, {"title": title, "detector": det, "details": {}})  # noqa: E731
    check("a new device (or a new open port) counts on Floor or Office, not on the Wiping VLAN, nor a randomized MAC",
          mt("New device on VLAN 10.69.0.0/16: aa:bb:cc:dd:ee:ff (Dell Inc.)", "arp_discovery")
          and mt("NEW open port 445/tcp (microsoft-ds) on 192.168.7.83", "kali_scan")
          and not mt("New device on VLAN 10.21.0.0/16: aa:bb:cc:dd:ee:ff (Dell Inc.)", "arp_discovery")
          and not mt("NEW open port 22/tcp (ssh) on 10.21.2.7", "kali_scan")
          and not mt("New device on VLAN 10.69.0.0/16: f6:11:7b:0a:85:75 (Unknown) — likely randomized Wi-Fi MAC",
                     "arp_discovery"))
    try:
        correlate._compile_rule({"id": "too-long", "name": "x", "severity": "medium", "type": "threshold", "count": 2,
                                 "window_minutes": 20000, "match": {"detector": "x"}}, set())
        too_long = False
    except ValueError:
        too_long = True
    check("...and a window longer than 7 days is still refused", too_long)

    # ---- (added) investigation graph -------------------------------------------
    group("graph")
    import soc_graph
    import soc_views
    check("an unknown incident has no graph", soc_graph.incident_graph("no-such-incident") is None)
    check("an unknown entity has no graph", soc_graph.entity_graph("ip:198.51.100.250") is None)
    if inc_id:
        g = soc_graph.incident_graph(inc_id)
        inc = correlate.get_incident(inc_id)
        check("the incident graph has an incident node and one alert node per alert",
              g is not None and sum(n["kind"] == "incident" for n in g["nodes"]) == 1
              and sum(n["kind"] == "alert" for n in g["nodes"]) >= len(inc["alert_ids"]))
        check("every alert in the incident is linked to it with a 'contains' edge",
              sum(1 for e in g["edges"] if e["label"] == "contains") == len(inc["alert_ids"]))
        check("node ids are stable/unique (no duplicate id across nodes)",
              len({n["id"] for n in g["nodes"]}) == len(g["nodes"]))
    top_entities = soc_views.entities(limit=1)
    if top_entities:
        ekey = top_entities[0]["key"]
        eg = soc_graph.entity_graph(ekey)
        check("the entity graph is centered on that entity and has at least one alert",
              eg is not None and any(n["id"] == ekey for n in eg["nodes"])
              and any(n["kind"] == "alert" for n in eg["nodes"]))
        check("every edge references a node that exists in the same graph",
              all(e["from"] in {n["id"] for n in eg["nodes"]} and e["to"] in {n["id"] for n in eg["nodes"]} for e in eg["edges"]))

    # ---- (added) suppressions created from the dashboard ------------------------
    group("dashboard suppressions")
    import suppression_admin
    tgt = emit_alert(Alert(type="port_scan", severity="medium", title="NEW open port 62078/tcp (tcpwrapped) on 10.1.1.10",
                           detector="kali_scan", source_ip="10.1.1.10", details={"port": 62078, "proto": "tcp"}), echo=False)
    other = emit_alert(Alert(type="port_scan", severity="medium", title="NEW open port 3389/tcp (ms-wbt-server) on 10.1.1.10",
                             detector="kali_scan", source_ip="10.1.1.10", details={"port": 3389, "proto": "tcp"}), echo=False)
    crit = emit_alert(Alert(type="intrusion", severity="critical", title="Something critical", detector="suricata",
                            source_ip="9.9.9.9"), echo=False)
    rule = suppression_admin.create_rule("selftest", "known benign Apple sync port", alert_id=tgt["id"], scope="similar")["rule"]
    check("a rule can be created from an alert (it always expires)", bool(rule["expires"]) and rule["added_by"] == "selftest")
    again = emit_alert(Alert(type="port_scan", severity="medium", title="NEW open port 62078/tcp (tcpwrapped) on 10.9.9.9",
                             detector="kali_scan", source_ip="10.9.9.9", details={"port": 62078, "proto": "tcp"}), echo=False)
    check("the same port on another host is now suppressed", again.get("suppressed_by") == rule["id"])
    other2 = emit_alert(Alert(type="port_scan", severity="medium", title="NEW open port 3389/tcp (ms-wbt-server) on 10.9.9.9",
                              detector="kali_scan", source_ip="10.9.9.9", details={"port": 3389, "proto": "tcp"}), echo=False)
    check("a different port (RDP) is NOT suppressed", not other2.get("suppressed_by"))
    check("critical alerts cannot be suppressed from the dashboard",
          _refused(lambda: suppression_admin.create_rule("selftest", "should be refused", alert_id=crit["id"])))
    check("a client-supplied regular expression is refused",
          _refused(lambda: suppression_admin.create_rule("selftest", "should be refused",
                                                          match={"detector": "aide", "title_regex": ".*"})))
    suppression_admin.delete_rule(rule["id"], "selftest")
    back = emit_alert(Alert(type="port_scan", severity="medium", title="NEW open port 62078/tcp (tcpwrapped) on 10.8.8.8",
                            detector="kali_scan", source_ip="10.8.8.8", details={"port": 62078, "proto": "tcp"}), echo=False)
    check("after deleting the rule the alert is shown again", not back.get("suppressed_by"))

    # ---- (added) playbooks ----------------------------------------------------
    group("playbooks")
    import playbooks
    ov = playbooks.overview()
    check("the playbook file is valid", not ov["errors"] and len(ov["playbooks"]) >= 1)
    check("the shipped playbooks are all disabled", not any(p["enabled"] for p in ov["playbooks"]))
    check("an incident alert triggers no action while disabled",
          playbooks.run_for_alert({"detector": "correlation", "severity": "critical", "title": "Incident #1",
                                   "details": {"incident_id": "x"}}) == [])

    # ---- (added) backup -------------------------------------------------------
    group("backup")
    bk = tmp / "bk"
    bp = subprocess.run(["bash", str(KALI / "backup_data.sh")], capture_output=True, text=True, timeout=120,
                        env=dict(os.environ, SOC_BACKUP_DIR=str(bk), SOC_DATA_DIR=str(tmp)))
    arch = sorted(bk.glob("soc-backup-*.tar.gz"))
    check("the backup script runs and writes an archive", bp.returncode == 0 and bool(arch), bp.stderr[-120:])
    if arch:
        names = subprocess.run(["tar", "-tzf", str(arch[-1])], capture_output=True, text=True).stdout
        check("the archive has the alerts and config but NOT the API keys",
              "data/alerts.json" in names and "config/" in names and "api_keys" not in names)

    # ---- severity model, grouping, test alerts, context ------------------------
    group("aide")
    import types
    import aide_to_alerts as A
    sev = A.severity_for
    check("AIDE: an ordinary path is normal", sev("/home/user/notes.txt") == "normal")
    check("AIDE: /etc, /root, /usr/local and .ssh paths are medium",
          all(sev(x) == "medium" for x in ("/etc/hosts", "/root/x", "/usr/local/bin/tool", "/home/u/.ssh/known_hosts")))
    check("AIDE: cron, spool cron and the project itself are critical",
          all(sev(x) == "critical" for x in ("/etc/cron.d/job", "/var/spool/cron/crontabs/root", "/opt/sentinel-soc/scripts/x.py")))

    def run_aide(paths, kind="Changed entries:"):
        report = kind + "\n---\n" + "".join(f"f =.... mc..H.. .  : {p_}\n" for p_ in paths) + "\nDetailed information:\n"
        real = A.subprocess.run
        A.subprocess.run = lambda *a_, **k_: types.SimpleNamespace(stdout=report, stderr="", returncode=4)
        try:
            A.run()
        finally:
            A.subprocess.run = real

    def suppressed_count() -> int:
        return len(soc_core.SUPPRESSED_LOG.read_text().splitlines()) if soc_core.SUPPRESSED_LOG.exists() else 0
    n = len(feed()); run_aide([f"/srv/data/pkg/file{i}" for i in range(7)])
    got = new_alerts(n)
    grp = find(got, "File integrity: 7 files changed under /srv/data/pkg")
    check("AIDE: more than 5 files under one folder become ONE alert", grp is not None and len(got) == 1)
    check("AIDE: the grouped alert lists the files and the count",
          grp is not None and grp["details"].get("count") == 7 and len(grp["details"].get("files", [])) == 7)
    n = len(feed()); run_aide([f"/srv/data/other/file{i}" for i in range(5)])
    check("AIDE: 5 files (not more) stay individual", len(new_alerts(n)) == 5)
    check("AIDE: a group is keyed by the parent directory, never the file itself",
          A.group_folder("/usr/bin/ac") == "/usr/bin" and A.group_folder("/etc/cron.daily/debsums") == "/etc/cron.daily"
          and A.group_folder("/srv/data/pkg/file0") == "/srv/data/pkg" and A.group_folder("/var/lib/a/b/c/d") == "/var/lib/a")
    n = len(feed()); run_aide([f"/usr/bin/newtool{i}" for i in range(7)], kind="Added entries:")
    got = new_alerts(n)
    check("AIDE: 7 new files directly in /usr/bin (a 3-level path) become ONE alert (regression: each was its own group, "
          "so a routine package install raised 34 separate critical alerts)",
          find(got, "7 files appeared under /usr/bin") is not None and len(got) == 1, str([a_["title"] for a_ in got][:3]))
    import suppression_admin as _sa
    q_rule = _sa.create_rule("selftest", "a folder whose files are known benign",
                             match={"detector": "aide", "title_contains": "/srv/data/quiet/"})["rule"]
    n = len(feed()); sup_before = suppressed_count()
    run_aide([f"/srv/data/quiet/file{i}" for i in range(8)])
    got, sup_new = new_alerts(n), suppressed_count() - sup_before
    check("AIDE: files a suppression rule matches stay individual (and are hidden, not grouped)",
          not got and sup_new == 8, f"feed +{len(got)}, suppressed +{sup_new}")
    _sa.delete_rule(q_rule["id"], "selftest")

    group("test alerts")
    calls = {"push": 0, "playbook": 0}
    import playbooks as _pb
    real_push, real_run = soc_core.notify_critical, _pb.run_for_alert
    soc_core.notify_critical = lambda *a_, **k_: calls.__setitem__("push", calls["push"] + 1)
    _pb.run_for_alert = lambda *a_, **k_: calls.__setitem__("playbook", calls["playbook"] + 1)
    try:
        def crit(**kw):
            return soc_core.emit_alert(soc_core.Alert(type="intrusion", severity="critical", title="selftest critical",
                                                      source_ip="10.9.9.9", detector=kw.pop("detector", "login_monitor"), **kw), echo=False)
        real_alert = crit()
        check("a normal alert has no 'test' key", "test" not in real_alert)
        check("a real critical alert does push and run playbooks (the probes work)", calls == {"push": 1, "playbook": 1}, str(calls))
        calls.update(push=0, playbook=0)
        by_flag = crit(test=True)
        by_detector = crit(detector="manual_test")
        os.environ["SOC_TEST_ALERTS"] = "1"
        try:
            by_env = crit()
        finally:
            del os.environ["SOC_TEST_ALERTS"]
        check("test=True, a test detector and SOC_TEST_ALERTS=1 all mark the record test: true",
              all(a_.get("test") is True for a_ in (by_flag, by_detector, by_env)))
        check("a test alert opens no incident", not any("incident_id" in a_["details"] for a_ in (by_flag, by_detector, by_env)))
        check("a test alert sends no push and runs no playbook", calls == {"push": 0, "playbook": 0}, str(calls))
    finally:
        soc_core.notify_critical, _pb.run_for_alert = real_push, real_run

    group("context")

    def emit_ctx(detector, source_ip, **details):
        return soc_core.emit_alert(soc_core.Alert(type="intrusion", severity="medium", title=f"selftest {detector}",
                                                  source_ip=source_ip, detector=detector, details=details), echo=False)
    a_pub = emit_ctx("zeek", "10.9.9.9", dst="8.8.8.8")
    check("a public IP in details.dst is checked for anonymizers",
          "anonymizer" in a_pub["details"] and a_pub["details"].get("anonymizer_ip") == "8.8.8.8")
    check("a public IP in details.ioc_ip is checked too",
          emit_ctx("malware_detector", "10.9.9.9", ioc_ip="1.1.1.1")["details"].get("anonymizer_ip") == "1.1.1.1")
    a_priv = emit_ctx("zeek", "10.9.9.9", dst="10.9.9.10")
    check("two private IPs get no anonymizer block", "anonymizer" not in a_priv["details"] and "anonymizer_ip" not in a_priv["details"])
    check("login_monitor's source is the actor", emit_ctx("login_monitor", "10.9.9.9")["details"].get("source_role") == "actor")
    check("kali_scan's source is the asset", emit_ctx("kali_scan", "10.9.9.9")["details"].get("source_role") == "asset")
    own = next(iter(soc_core.own_ips()), None)
    if own:
        check("source_is_self is set for this box's own address", emit_ctx("kali_scan", own)["details"].get("source_is_self") is True)
    check("source_is_self is absent for another machine", "source_is_self" not in emit_ctx("kali_scan", "10.9.9.9")["details"])

    group("nmap grouping")

    def host_xml(ip, ports):
        body = "".join(f'<port protocol="tcp" portid="{p_}"><state state="open"/><service name="http" method="probed" conf="10"/></port>' for p_ in ports)
        return ('<?xml version="1.0"?><nmaprun><host><status state="up"/>'
                f'<address addr="{ip}" addrtype="ipv4"/><hostnames/><ports>{body}</ports></host></nmaprun>')
    st2 = tmp / "port_state_group.json"
    (tmp / "g1.xml").write_text(host_xml("10.1.2.6", [22])); N.run(tmp / "g1.xml", {22}, st2)
    n = len(feed())
    (tmp / "g2.xml").write_text(host_xml("10.1.2.6", [22, 8001, 8002, 8003, 8004, 8005])); N.run(tmp / "g2.xml", {22}, st2)
    got = new_alerts(n)
    check("more than 4 new ports on one host become ONE alert", find(got, "NEW open ports on 10.1.2.6") is not None
          and not any(a_["title"].startswith("NEW open port ") for a_ in got), str([a_["title"] for a_ in got]))
    (tmp / "g3.xml").write_text(host_xml("10.1.2.7", [22])); N.run(tmp / "g3.xml", {22}, st2)
    n = len(feed())
    (tmp / "g4.xml").write_text(host_xml("10.1.2.7", [22, 8001, 8002, 8003, 8004])); N.run(tmp / "g4.xml", {22}, st2)
    got = new_alerts(n)
    check("4 new ports stay individual", sum(a_["title"].startswith("NEW open port ") for a_ in got) == 4, str([a_["title"] for a_ in got]))

    # ---- (added) what the scan already learns: SMB signing, Windows builds ---------------------------
    group("scan facts")
    import host_facts
    from datetime import date
    today_ = date(2026, 9, 23)
    ws = lambda v, d=today_: host_facts.windows_support(v, d)
    check("a Windows build whose support ended years ago is medium", (ws("6.1.7601") or {}).get("severity") == "medium"
          and (ws("10.0.10240") or {}).get("severity") == "medium")
    w10 = ws("10.0.19041")
    check("build 19041 is named as a FAMILY (RDP cannot tell 2004 from 22H2), dated by its last mainstream edition",
          w10 and "19041 family" in w10["label"] and w10["ended"] == "2025-10-14" and w10["severity"] == "normal"
          and (host_facts.windows_support("10.0.19041", date(2026, 10, 15)) or {}).get("severity") == "medium")
    w11 = ws("10.0.22621")
    check("build 22621 (Windows 11 22H2 or 23H2) is flagged as a family, dated by 23H2 Home/Pro, and says Enterprise/Education "
          "may still be covered (NTLM cannot tell the edition)",
          w11 and "22621 family" in w11["label"] and w11["ended"] == "2025-11-11" and w11["severity"] == "normal"
          and "2026-11-10" in w11["caveat"])
    check("a build with no caveat has an empty one", ws("6.1.7601")["caveat"] == "")
    check("one that ended within the last year is only normal, and becomes medium after a year",
          (ws("10.0.19045") or {}).get("severity") == "normal"
          and (host_facts.windows_support("10.0.19045", date(2026, 10, 15)) or {}).get("severity") == "medium")
    check("a supported build, one shared with a supported Server release, and a build we do not know raise nothing",
          ws("10.0.26100") is None and ws("10.0.17763") is None and ws("10.0.22631") is None and ws("garbage") is None
          and ws(None) is None)
    check("SMB signing is read from the script output",
          host_facts.parse_smb_signing("3.1.1: Message signing enabled and required") == "required"
          and host_facts.parse_smb_signing("3.1.1: Message signing enabled but not required") == "not_required"
          and host_facts.parse_smb_signing("nothing useful") is None)

    facts_xml = tmp / "facts_scan.xml"
    facts_xml.write_text("""<?xml version="1.0"?><nmaprun><host><status state="up"/>
<address addr="10.69.9.1" addrtype="ipv4"/><hostnames/>
<ports><port protocol="tcp" portid="3389"><state state="open"/><service name="ms-wbt-server" method="probed" conf="10"/>
<script id="rdp-ntlm-info" output="x"><elem key="Product_Version">6.1.7601</elem><elem key="NetBIOS_Computer_Name">OLDPC</elem>
<elem key="DNS_Domain_Name">era.local</elem></script></port></ports>
<hostscript><script id="smb2-security-mode" output="&#10;  3.1.1: &#10;    Message signing enabled but not required"/>
<script id="nbstat" output="NetBIOS name: OLDPC, NetBIOS user: &lt;unknown&gt;"/></hostscript>
<os><osmatch name="Microsoft Windows 10 2004" accuracy="96"/></os></host></nmaprun>""")
    ex = host_facts.extract(facts_xml)
    f1 = ex.get("10.69.9.1", {})
    check("the OS guess, Windows build, names and SMB signing are read from the scan's XML",
          f1.get("os") == {"name": "Microsoft Windows 10 2004", "accuracy": 96}
          and f1.get("windows", {}).get("Product_Version") == "6.1.7601"
          and f1.get("smb_signing") == "not_required" and f1.get("netbios_name") == "OLDPC", str(f1))
    host_facts.TRUST_NTLM_BUILD = False
    fnd_off = host_facts.findings(ex, {}, today_)
    host_facts.TRUST_NTLM_BUILD = True
    check("the kill switch (SOC_TRUST_NTLM_BUILD=0) turns the end-of-support finding off and leaves SMB signing on",
          [x_["title"].split(" on ")[0] for x_ in fnd_off] == ["SMB signing not required"], str([x_["title"] for x_ in fnd_off]))
    check("the build is kept in the extracted facts either way", ex["10.69.9.1"]["windows"]["Product_Version"] == "6.1.7601")
    check("by default the finding is on", host_facts.TRUST_NTLM_BUILD is True)
    fnd = host_facts.findings(ex, {}, today_)
    check("that host yields exactly two findings: SMB signing (normal) and unsupported Windows (medium)",
          sorted((x_["title"].split(" on ")[0], x_["severity"]) for x_ in fnd)
          == [("SMB signing not required", "normal"), ("Unsupported Windows", "medium")], str([x_["title"] for x_ in fnd]))

    n = len(feed())
    N.run(facts_xml, {22}, tmp / "scanfacts_state.json")
    got = new_alerts(n)
    a_smb = find(got, "SMB signing not required on 10.69.9.1")
    a_win = find(got, "Unsupported Windows on 10.69.9.1")
    check("the scan import raises both findings, with the computer name from RDP",
          a_smb is not None and a_win is not None and a_win["hostname"] == "OLDPC" and a_win["severity"] == "medium")
    check("...and tags SMB signing MITRE T1557.001 (relay exposure), but not the Windows finding (no honest mapping)",
          "T1557.001" in tags(a_smb) and not tags(a_win))
    n = len(feed())
    N.run(facts_xml, {22}, tmp / "scanfacts_state.json")
    check("...and does not raise them again on the next scan (regression guard: findings alert once)",
          find(new_alerts(n), "SMB signing not required") is None and find(new_alerts(n), "Unsupported Windows") is None)
    # a PC that Active Directory already covers: ad_inventory raises the (more exact) alert, so no second one here
    from datetime import datetime as _dt, timedelta as _td, timezone as _tz
    _now = _dt.now(_tz.utc)
    (tmp / "ad_inventory.json").write_text(json.dumps({"computers": {
        "ADPC": {"enabled": True, "last_logon": (_now - _td(days=3)).isoformat()},
        "OFFPC": {"enabled": False, "last_logon": (_now - _td(days=3)).isoformat()},
        "STALEPC": {"enabled": True, "last_logon": (_now - _td(days=200)).isoformat()},
        "NOLOGON": {"enabled": True, "last_logon": None},
        "BADDATE": {"enabled": True, "last_logon": "not-a-date"}}}))
    check("only enabled AD computers that signed in within 90 days count as covered (names upper-case)",
          host_facts.ad_covered_names() == {"ADPC"}, str(host_facts.ad_covered_names()))
    check("a missing or unreadable inventory covers nothing", host_facts.ad_covered_names(tmp / "nope.json") == set())
    ad_xml = tmp / "ad_scan.xml"
    ad_xml.write_text(facts_xml.read_text().replace("10.69.9.1", "10.69.9.2").replace("OLDPC", "adpc"))
    n = len(feed())
    N.run(ad_xml, {22}, tmp / "scanfacts_state.json")
    got = new_alerts(n)
    check("for an AD-covered PC the SMB-signing alert still fires but the NTLM 'Unsupported Windows' alert does not",
          find(got, "SMB signing not required on 10.69.9.2") is not None and find(got, "Unsupported Windows on 10.69.9.2") is None)
    check("...yet the finding is tracked, so it can never later show up as 'No longer detected'",
          "10.69.9.2:windows-support" in json.loads(N._vuln_state_path(tmp / "scanfacts_state.json").read_text())["hosts"])
    (tmp / "ad_inventory.json").unlink()
    sf = soc_core.load_assets()["3c:bb:cc:00:00:01"].get("scan_facts", {})
    check("what the scan learned is kept on that host's asset record",
          sf.get("windows", {}).get("Product_Version") == "6.1.7601" and sf.get("smb_signing") == "not_required"
          and sf.get("os", {}).get("name") == "Microsoft Windows 10 2004", str(sf))

    check("record_asset_scan_facts never creates an asset, and an identical repeat changes nothing",
          soc_core.record_asset_scan_facts([{"mac": "00:00:00:00:00:98", "facts": {"os": {"name": "x"}}}]) == 0
          and "00:00:00:00:00:98" not in soc_core.load_assets()
          and soc_core.record_asset_scan_facts([{"mac": "3c:bb:cc:00:00:01", "facts": {"smb_signing": "not_required"}}]) == 0)
    soc_core.record_asset_scan_facts([{"mac": "3c:bb:cc:00:00:01", "facts": {"smb_signing": "required"}}])
    sf = soc_core.load_assets()["3c:bb:cc:00:00:01"]["scan_facts"]
    check("a fact that is present replaces the old value, and facts missing from a later scan are kept",
          sf["smb_signing"] == "required" and sf["windows"]["Product_Version"] == "6.1.7601")

    # ---- (added) Active Directory inventory (no LDAP here: the pure parts) ----------------------------
    group("ad inventory")
    import ad_inventory as AD
    d0 = date(2026, 9, 23)
    sup = lambda os_, v, d=d0: AD.os_support(os_, v, d)
    check("AD's operatingSystemVersion is parsed to a build",
          AD.parse_build("10.0 (26100)") == 26100 and AD.parse_build("6.3 (9600)") == 9600 and AD.parse_build(None) is None)
    check("Windows 10 22H2 is unsupported whatever the edition; 24H2 Pro is ending soon; 25H2 is supported",
          sup("Windows 10 Pro", "10.0 (19045)")["status"] == "unsupported"
          and sup("Windows 10 Enterprise", "10.0 (19045)")["status"] == "unsupported"
          and sup("Windows 11 Pro", "10.0 (26100)")["status"] == "ending_soon"
          and sup("Windows 11 Pro", "10.0 (26200)")["status"] == "supported")
    check("the edition picks the date: 23H2 Pro ended, 23H2 Enterprise still supported; Pro Education follows Pro",
          sup("Windows 11 Pro", "10.0 (22631)")["status"] == "unsupported"
          and sup("Windows 11 Enterprise", "10.0 (22631)")["ends"] == "2026-11-10"
          and sup("Windows 11 Pro Education", "10.0 (22631)")["track"] == "Home/Pro"
          and sup("Windows 11 Education", "10.0 (22631)")["track"] == "Enterprise/Education")
    check("servers use their own table; LTSC, non-Windows and unknown builds are left alone",
          sup("Windows Server 2019 Standard", "10.0 (17763)")["status"] == "supported"
          and sup("Windows Server 2012 R2 Standard", "6.3 (9600)")["status"] == "unsupported"
          and sup("Windows 10 Enterprise LTSC", "10.0 (17763)") is None and sup("Linux", "5.4") is None
          and sup("Windows 11 Pro", "10.0 (99999)") is None and sup(None, None) is None)
    check("the OU path is read from the DN", AD._ou("CN=PC1,OU=Accounting,OU=Calgary,DC=era,DC=local") == "Calgary/Accounting")

    def comp(name, os_, ver, last="2026-09-20", enabled=True):
        return {"name": name, "dns": f"{name.lower()}.era.local", "os": os_, "os_version": ver, "build": AD.parse_build(ver),
                "enabled": enabled, "last_logon": f"{last}T00:00:00+00:00" if last else None,
                "created": "2024-01-01T00:00:00+00:00", "ou": "Calgary/Accounting", "support": AD.os_support(os_, ver, d0)}

    def usr(sam, last="2026-09-20", enabled=True):
        return {"sam": sam, "display": sam.title(), "enabled": enabled,
                "last_logon": f"{last}T00:00:00+00:00" if last else None, "created": "2024-01-01T00:00:00+00:00", "ou": "Calgary"}

    snap1 = {"generated": "2026-09-23T12:00:00+00:00",
             "computers": {"OLD10": comp("OLD10", "Windows 10 Pro", "10.0 (19045)"),
                           "NEW24": comp("NEW24", "Windows 11 Pro", "10.0 (26100)"),
                           "GONE": comp("GONE", "Windows 10 Pro", "10.0 (19045)", last="2025-01-01"),
                           "AZUREADSSOACC": comp("AZUREADSSOACC", None, None, last=None)},
             "users": {"ana": usr("ana"), "lyndsay": usr("lyndsay", last="2025-12-01"), "krbtgt": usr("krbtgt", last=None)},
             "privileged": {"Domain Admins": ["administrator", "ana"]}}
    al1, st1 = AD.evaluate(snap1, {}, {"OLD10": "10.69.1.10", "ROGUE-PC": "10.69.1.99", "TEST-WIN10": "10.69.250.11"}, d0,
                           ["TEST-*"])
    t1 = [x_["title"] for x_ in al1]
    check("first run: one baseline alert for the privileged members, and no 'new computer/user' flood",
          "AD privileged group members recorded (baseline)" in t1 and not any(x_.startswith("New ") for x_ in t1), str(t1))
    w = next((x_ for x_ in al1 if x_["title"].startswith("Unsupported Windows (per AD) on OLD10")), None)
    check("an unsupported active PC alerts with its IP from the scan; a stale one does not (it is in the stale list instead)",
          w is not None and w["source_ip"] == "10.69.1.10" and w["type"] == "vuln" and w["severity"] == "medium"
          and not any("on GONE" in x_ for x_ in t1), str(t1))
    check("a release losing support within the warning window is one grouped alert",
          any(x_.startswith("1 computer(s) lose Windows support on 2026-10-13") for x_ in t1), str(t1))
    check("stale computers and users are summarised, never-sign-in built-ins excluded",
          st1["stale_computers"] == ["GONE"] and st1["stale_users"] == ["lyndsay"], str((st1["stale_computers"], st1["stale_users"])))
    check("a Windows machine on the network that AD does not know alerts (medium); an ignored pattern does not",
          any(x_ == "Windows host not in the domain: ROGUE-PC (10.69.1.99)" for x_ in t1)
          and not any("TEST-WIN10" in x_ for x_ in t1))

    snap2 = json.loads(json.dumps(snap1))
    snap2["privileged"]["Domain Admins"] = ["administrator", "mallory"]
    snap2["computers"]["LAPTOP9"] = comp("LAPTOP9", "Windows 11 Pro", "10.0 (26200)")
    snap2["users"]["mallory"] = usr("mallory")
    al2, st2 = AD.evaluate(snap2, st1, {"OLD10": "10.69.1.10", "ROGUE-PC": "10.69.1.99"}, d0, ["TEST-*"])
    by = {x_["title"]: x_ for x_ in al2}
    check("next day: an addition to Domain Admins is critical, a removal is medium",
          by.get("Added to Domain Admins: mallory", {}).get("severity") == "critical"
          and by.get("Removed from Domain Admins: ana", {}).get("severity") == "medium", str(list(by)))
    check("...new computer and user accounts are reported once each",
          "New computer in AD: LAPTOP9" in by and "New user account in AD: mallory (Mallory)" in by
          and by["New computer in AD: LAPTOP9"]["severity"] == "medium", str(list(by)))
    check("...and nothing already reported comes back (unsupported PC, ending-soon group, stale lists, not-in-domain host)",
          not any(x_.startswith(("Unsupported Windows", "1 computer(s) lose", "Windows host not in")) or "unused for" in x_
                  for x_ in by), str(list(by)))
    snap3 = json.loads(json.dumps(snap2))
    snap3["computers"]["OLD10"] = comp("OLD10", "Windows 11 Pro", "10.0 (26200)")
    al3, st3 = AD.evaluate(snap3, st2, {}, d0)
    check("an upgraded PC leaves the unsupported list quietly", "OLD10" not in st3["unsupported"] and not al3, str(al3))
    # sign-ins that should not happen: an account on leave, a dormant account (unused for STALE_DAYS)
    sa = json.loads(json.dumps(snap3))
    sa["users"]["dormant"] = usr("dormant", last="2025-01-10")
    sa["users"]["oldadmin"] = usr("oldadmin", last="2025-02-01")
    sa["privileged"]["Domain Admins"] = sa["privileged"]["Domain Admins"] + ["oldadmin"]
    _, sta = AD.evaluate(sa, st3, {}, d0)
    sb = json.loads(json.dumps(sa))
    sb["users"]["lyndsay"]["last_logon"] = "2026-09-22T08:00:00+00:00"   # on leave, yet a new sign-in
    sb["users"]["dormant"]["last_logon"] = "2026-09-21T08:00:00+00:00"   # 600 days without use, then used
    sb["users"]["oldadmin"]["last_logon"] = "2026-09-21T09:00:00+00:00"
    sb["users"]["ana"]["last_logon"] = "2026-09-23T08:00:00+00:00"       # a normal, active user signing in again
    alb, stb = AD.evaluate(sb, sta, {}, d0, on_leave=["Lyndsay"])
    byb = {x_["title"].split(":")[0]: [] for x_ in alb}
    for x_ in alb:
        byb[x_["title"].split(":")[0]].append((x_["user"], x_["severity"]))
    check("a sign-in by an account on the on-leave watchlist is critical",
          byb.get("Sign-in by an account on leave") == [("lyndsay", "critical")], str(byb))
    check("a dormant account used again is medium, critical when it is privileged; an active user signing in is nothing",
          sorted(byb.get("Dormant account used again", [])) == [("dormant", "medium"), ("oldadmin", "critical")]
          and not any(u_ == "ana" for v_ in byb.values() for u_, _ in v_), str(byb))
    alc, _ = AD.evaluate(sb, stb, {}, d0, on_leave=["lyndsay"])
    check("...each sign-in alerts once", not alc, str([x_["title"] for x_ in alc]))
    _, st_first = AD.evaluate(sb, {k_: v_ for k_, v_ in sta.items() if k_ != "last_logons"}, {}, d0, on_leave=["lyndsay"])
    check("...and the first run after the upgrade only records the dates",
          "last_logons" in st_first and not [x_ for x_ in AD.evaluate(sb, {k_: v_ for k_, v_ in sta.items()
                                                                            if k_ != "last_logons"}, {}, d0,
                                                                           on_leave=["lyndsay"])[0]
                                             if "Sign-in" in x_["title"] or "Dormant" in x_["title"]])
    many = json.loads(json.dumps(snap3))
    for i_ in range(AD.MAX_INDIVIDUAL_NEW + 5):
        many["computers"][f"BULK{i_}"] = comp(f"BULK{i_}", "Windows 11 Pro", "10.0 (26200)")
    al4, _ = AD.evaluate(many, st3, {}, d0)
    check("a burst of new objects (re-baseline) is one summary alert, not one per object",
          [x_["title"] for x_ in al4] == [f"{AD.MAX_INDIVIDUAL_NEW + 5} new computers in AD"], str([x_["title"] for x_ in al4]))

    AD.STATE_FILE = tmp / "ad_inventory_state.json"
    n = len(feed())
    AD._failure({}, "ConnectionError: 10.69.0.14: bind refused", False)
    AD._failure(json.loads(AD.STATE_FILE.read_text()), "ConnectionError: again", False)
    check("a failed AD read alerts once, not on every retry",
          sum(a_["title"] == "AD inventory cannot read Active Directory" for a_ in new_alerts(n)) == 1)

    dc_ = AD.daily_counts(snap1, st1, d0)
    check("the daily history line counts active/stale computers, stale users, admins and unsupported PCs",
          dc_["computers"]["stale"] == 1 and dc_["users"]["stale"] == 1 and dc_["privileged_accounts"] == 2
          and dc_["os_unsupported"] == 1 and dc_["os_ending_soon"] == 1 and dc_["not_in_domain"] == 1, str(dc_))
    AD.HISTORY_FILE = tmp / "ad_history.jsonl"
    AD.record_history(dc_)
    AD.record_history({**dc_, "not_in_domain": 5})
    AD.record_history({**dc_, "date": "2026-09-24"})
    hist = [json.loads(l_) for l_ in AD.HISTORY_FILE.read_text().splitlines()]
    check("one history line per day: a re-run the same day replaces it",
          [(h_["date"], h_["not_in_domain"]) for h_ in hist] == [("2026-09-23", 5), ("2026-09-24", 1)], str(hist))

    pc = AD.privileged_changes({"Domain Admins": ["administrator"]},
                               {"Domain Admins": ["administrator", "eve"], "DnsAdmins": ["bob"]})
    check("the 15-minute privileged check: an addition is critical, a group seen for the first time is only recorded",
          [(x_["title"], x_["severity"]) for x_ in pc] == [("Added to Domain Admins: eve", "critical")], str(pc))

    import ad_risks as AR

    def acct(sam, uac=0x200, spn=(), pwd="2026-01-01", last="2026-09-20", computer=False, os_=None, laps=False):
        return {"sam": sam, "computer": computer, "uac": uac, "spn": list(spn), "pwd_last_set": pwd, "last_logon": last,
                "admin_count": 0, "os": os_, "laps": laps}
    rdata = {"accounts": [
        acct("svc_sql", spn=["MSSQLSvc/db1:1433"]), acct("svc_admin", spn=["HTTP/app"], pwd="2020-01-01"),
        acct("krbtgt", uac=0x202, spn=["kadmin/changepw"], pwd="2025-03-19"),
        acct("nopreauth", uac=0x200 | AR.DONT_REQ_PREAUTH), acct("empty", uac=0x200 | AR.PASSWD_NOTREQD),
        acct("olddisabled", uac=0x202 | AR.DONT_REQ_PREAUTH),
        acct("DC1$", uac=AR.SERVER_TRUST_ACCOUNT | AR.TRUSTED_FOR_DELEGATION, computer=True, os_="Windows Server 2019"),
        acct("APP1$", uac=0x1000 | AR.TRUSTED_FOR_DELEGATION, computer=True, os_="Windows Server 2019"),
        acct("PC1$", uac=0x1000, computer=True, os_="Windows 11 Pro", laps=True),
        acct("PC2$", uac=0x1000, computer=True, os_="Windows 11 Pro")],
        "policy": {"min_length": 7, "lockout_threshold": 0, "max_age_days": 42, "history": 24}, "laps_in_schema": True}
    rf = {x_["id"]: x_ for x_ in AR.risk_findings(rdata, ["svc_admin"], d0)}
    check("Kerberoastable accounts are found, split by privilege (critical when an admin), krbtgt excluded",
          rf.get("kerberoastable_privileged", {}).get("accounts") == ["svc_admin"]
          and rf["kerberoastable_privileged"]["severity"] == "critical"
          and rf.get("kerberoastable", {}).get("accounts") == ["svc_sql"], str(list(rf)))
    check("AS-REP roasting and empty-password flags are found on enabled accounts only",
          rf.get("asrep_roastable", {}).get("accounts") == ["nopreauth"]
          and rf.get("password_not_required", {}).get("accounts") == ["empty"])
    check("unconstrained delegation on a server is flagged, on a domain controller it is not",
          rf.get("unconstrained_delegation", {}).get("accounts") == ["APP1$"])
    check("old admin password, old krbtgt, a short minimum length and no lockout are flagged",
          rf.get("privileged_old_password", {}).get("accounts") == ["svc_admin"] and "krbtgt_old_password" in rf
          and "weak_password_policy" in rf and "no_lockout" in rf)
    check("LAPS coverage counts active workstations only (servers and DCs left out)",
          rf.get("laps_missing", {}).get("accounts") == ["PC2"] and rf["laps_missing"]["severity"] == "normal"
          and "1 of 2" in rf["laps_missing"]["title"], str(rf.get("laps_missing")))
    rdata["policy"] = {"min_length": 14, "lockout_threshold": 10}
    check("a sound policy raises nothing", not {"weak_password_policy", "no_lockout"} & {
        x_["id"] for x_ in AR.risk_findings(rdata, [], d0)})

    snapr = json.loads(json.dumps(snap3))
    snapr["risks"] = [{"id": "kerberoastable", "severity": "medium", "title": "Kerberoastable account(s) in AD",
                       "description": "d", "accounts": ["svc_sql"], "mitre": None}]
    al5, st5 = AD.evaluate(snapr, st3, {}, d0)
    snapr["risks"][0]["accounts"] = ["svc_sql", "svc_web"]
    al6, st6 = AD.evaluate(snapr, st5, {}, d0)
    snapr["risks"] = []
    al7, _ = AD.evaluate(snapr, st6, {}, d0)
    check("a weakness alerts when it appears, again only when it gains accounts, and a note when it is fixed",
          [x_["title"] for x_ in al5] == ["Kerberoastable account(s) in AD: svc_sql"]
          and len(al6) == 1 and "svc_web" in al6[0]["description"]
          and [x_["title"] for x_ in al7] == ["AD weakness fixed: kerberoastable"],
          str(([x_["title"] for x_ in al5], [x_["title"] for x_ in al6], [x_["title"] for x_ in al7])))
    import mitre_tags
    check("AD findings carry MITRE tags (Domain Admin addition T1098, Kerberoasting T1558.003)",
          [t_["technique"] for t_ in mitre_tags.tag({"detector": "ad_inventory", "title": "Added to Domain Admins: eve"})]
          == ["T1098"]
          and [t_["technique"] for t_ in mitre_tags.tag({"detector": "ad_inventory",
                                                          "title": "Kerberoastable account(s) in AD: svc_sql"})]
          == ["T1558.003"])

    # identity on every alert (ad_identity.py via emit_alert); OLDPC is 10.69.9.1 in the scan-facts test above
    (tmp / "ad_inventory.json").write_text(json.dumps({
        "generated": "2026-09-23T12:00:00+00:00",
        "computers": {"OLDPC": comp("OLDPC", "Windows 10 Pro", "10.0 (19045)")},
        "users": {"jdoe": usr("jdoe")}, "privileged": {"Domain Admins": ["jdoe"]}}))
    n = len(feed())
    soc_core.emit_alert(soc_core.Alert(type="vuln", severity="normal", title="selftest identity by IP",
                                       source_ip="10.69.9.1", detector="kali_scan"), echo=False)
    soc_core.emit_alert(soc_core.Alert(type="intrusion", severity="normal", title="selftest identity by user",
                                       user="ERA\\JDoe", detector="suricata"), echo=False)
    soc_core.emit_alert(soc_core.Alert(type="intrusion", severity="normal", title="selftest local user",
                                       user="jdoe", detector="login_monitor"), echo=False)
    got = new_alerts(n)
    i1 = (find(got, "selftest identity by IP") or {}).get("details", {}).get("identity", {})
    i2 = (find(got, "selftest identity by user") or {}).get("details", {}).get("identity", {})
    check("an alert about an IP gets the AD computer behind it (via the scan's NetBIOS name), with its OU and OS support",
          (i1.get("computers") or [{}])[0].get("name") == "OLDPC" and i1["computers"][0]["ou"] == "Calgary/Accounting"
          and i1["computers"][0]["os_support"] == "unsupported" and i1["computers"][0]["ip"] == "10.69.9.1", str(i1))
    check("an alert about DOMAIN\\user gets the AD user, with its privileged groups",
          (i2.get("users") or [{}])[0].get("sam") == "jdoe"
          and i2["users"][0]["privileged_groups"] == ["Domain Admins"], str(i2))
    check("a detector watching this appliance does not map its local users to AD",
          "identity" not in (find(got, "selftest local user") or {}).get("details", {}))

    # the monthly AD review workbook (ad_report.py), from the test inventory above (OLDPC, jdoe)
    import ad_report
    import openpyxl
    rep = tmp / "ad_review.xlsx"
    counts_ = ad_report.build(rep, d0)
    wb_ = openpyxl.load_workbook(rep)
    urow = {r_[0]: r_ for r_ in wb_["Users"].iter_rows(min_row=2, values_only=True)}
    check("the monthly AD review workbook has its sheets, and a Domain Admin is marked for review, not cleanup",
          wb_.sheetnames == ["Summary", "Computers", "Users", "Admins", "Domain weaknesses"]
          and counts_["computers"] == 1 and urow.get("jdoe", (None, None, None))[2] == "review", str(urow)[:200])

    # changes to GPOs, links, trusts and domain permissions (ad_changes.py; pure comparison, no LDAP)
    import ad_changes as ACH
    check("gPLink is parsed, and a disabled link is marked",
          ACH.parse_gplink("[LDAP://cn={aaa-1},cn=policies,cn=system,DC=x;0][LDAP://cn={bbb-2},cn=policies,cn=system,DC=x;1]")
          == ["{AAA-1}", "!{BBB-2}"])
    dom = "S-1-5-21-1-2-3"
    aces0 = [f"ACCESS_ALLOWED_OBJECT|{dom}-516|0x100|{ACH.GET_CHANGES_ALL}", f"ACCESS_ALLOWED|{dom}-512|0xf01ff|",
             "ACCESS_ALLOWED|S-1-5-11|0x20094|"]
    check("DCSync holders: the extended right or full control counts, plain read does not",
          ACH.dcsync_holders(aces0) == [f"{dom}-512", f"{dom}-516"])
    base0 = {"gpos": {"{G1}": {"name": "Accounting Local Admin", "version": 5, "flags": 0},
                      "{G2}": {"name": "Default Domain Controllers Policy", "version": 2, "flags": 0}},
             "links": {"OU=Accounting,DC=x": ["{G1}"]}, "trusts": {},
             "acl": {"adminsdholder": ["ACCESS_ALLOWED|S-1-5-18|0xf01ff|"], "domain": aces0},
             "dcsync": ACH.dcsync_holders(aces0)}
    first = ACH.change_alerts(None, base0, {})
    check("first reading: one baseline alert, nothing else",
          [x_["title"] for x_ in first] == ["AD change monitoring started (baseline)"])
    check("...and an identical second reading raises nothing", ACH.change_alerts(base0, json.loads(json.dumps(base0)), {}) == [])
    now1 = json.loads(json.dumps(base0))
    now1["gpos"]["{G1}"]["version"] = 6
    now1["gpos"]["{G2}"]["version"] = 3
    now1["gpos"]["{G3}"] = {"name": "Evil", "version": 1, "flags": 0}
    now1["links"]["DC=x"] = ["{G3}"]
    now1["links"]["OU=Accounting,DC=x"] = ["!{G1}"]
    now1["trusts"]["partner.example"] = {"direction": 3, "type": 2, "attributes": 8}
    now1["acl"]["adminsdholder"].append("ACCESS_ALLOWED|S-1-5-21-1-2-3-1105|0xf01ff|")
    now1["acl"]["domain"].append(f"ACCESS_ALLOWED_OBJECT|{dom}-1105|0x100|{ACH.GET_CHANGES_ALL}")
    now1["dcsync"] = ACH.dcsync_holders(now1["acl"]["domain"])
    ch = {x_["title"]: x_ for x_ in ACH.change_alerts(base0, now1, {f"{dom}-1105": "mallory"})}
    check("a GPO edit is medium, an edit of the DC policy critical, a new GPO and its link reported",
          ch.get("GPO modified: Accounting Local Admin", {}).get("severity") == "medium"
          and ch.get("GPO modified: Default Domain Controllers Policy", {}).get("severity") == "critical"
          and "New GPO created: Evil" in ch and "GPO linked: Evil on DC=x" in ch, str(list(ch)))
    check("a link being disabled is reported as such, not as unlinked + linked",
          "GPO link disabled: Accounting Local Admin on OU=Accounting,DC=x" in ch
          and not any(t_.startswith("GPO unlinked") for t_ in ch), str(list(ch)))
    check("a new trust, an AdminSDHolder change and new DCSync rights are critical, named by account",
          ch.get("New domain trust: partner.example", {}).get("severity") == "critical"
          and ch.get("AdminSDHolder permissions changed", {}).get("severity") == "critical"
          and ch.get("DCSync rights granted: mallory", {}).get("severity") == "critical", str(list(ch)))
    check("GPO and DCSync alerts carry MITRE T1484.001 and T1003.006",
          [t_["technique"] for t_ in mitre_tags.tag({"detector": "ad_inventory", "title": "GPO modified: X"})] == ["T1484.001"]
          and [t_["technique"] for t_ in mitre_tags.tag({"detector": "ad_inventory",
                                                          "title": "DCSync rights granted: mallory"})] == ["T1003.006"])

    # ---- (added) watchlists ---------------------------------------------------
    group("watchlists")
    import watchlists as W
    names = {w["name"] for w in W.list_watchlists()}
    check("all nine watchlists are registered", names == {"trusted_ips", "bad_ips", "bad_domains", "bad_hashes",
                                                          "dhcp_servers", "ra_sources", "sensitive_vlans", "untrusted_vlans",
                                                          "on_leave_accounts"}, str(names))
    before = W.get("trusted_ips")["count"]
    W.add("trusted_ips", "203.0.113.9/32", "selftest")
    check("adding a valid CIDR to a watchlist works", W.get("trusted_ips")["count"] == before + 1)
    try:
        W.add("trusted_ips", "203.0.113.9/32", "selftest")
        check("adding the same entry twice is refused", False)
    except ValueError:
        check("adding the same entry twice is refused", True)
    try:
        W.add("trusted_ips", "not-an-ip", "selftest")
        check("an invalid entry is refused (per-list validation)", False)
    except ValueError:
        check("an invalid entry is refused (per-list validation)", True)
    W.remove("trusted_ips", "203.0.113.9/32", "selftest")
    check("removing an entry works", W.get("trusted_ips")["count"] == before)
    try:
        W.get("does-not-exist")
        check("an unknown watchlist name is refused", False)
    except ValueError:
        check("an unknown watchlist name is refused", True)
    log = [json.loads(ln) for ln in soc_core.DATA_DIR.joinpath("watchlist_log.jsonl").read_text().splitlines()]
    check("the add and remove were audited", sum(1 for e in log if e["entry"] == "203.0.113.9/32") == 2)
    real_bad_domains = KALI / "bad_domains.txt"
    real_before = real_bad_domains.read_text() if real_bad_domains.exists() else ""
    W.add("bad_domains", "totally-fake-selftest-domain.example", "selftest")
    check("bad_ips/bad_domains (kali/-backed lists) are isolated too, not the real project files",
          "totally-fake-selftest-domain.example" in W.get("bad_domains")["entries"]
          and (not real_bad_domains.exists() or real_bad_domains.read_text() == real_before))

    # ---- (added) broadcast-level attacks: rogue DHCP server, rogue IPv6 router, LLMNR/NBT-NS poisoner, ARP owner ----
    group("l2 detections")
    import l2_watch
    import poisoner_canary as PC
    import socket as _socket
    import struct as _struct
    import threading as _threading
    for ip_ in ("10.0.0.51", "10.0.0.52"):
        W.add("dhcp_servers", ip_, "selftest")
    W.add("ra_sources", "fe80::1", "selftest")

    dw = l2_watch.DhcpWatch()
    drow = lambda server, msgs: {"msg_types": msgs, "server_addr": server, "mac": "aa:bb:cc:00:00:99", "assigned_addr": "10.0.0.77"}
    n = len(feed())
    check("an OFFER from a trusted DHCP server is ignored", dw.feed(drow("10.0.0.51", ["OFFER"])) is False and not new_alerts(n))
    check("a client's DISCOVER (no server in it) is ignored", dw.feed(drow(None, ["DISCOVER"])) is False)
    check("an ACK from an unknown DHCP server is a critical alert, tagged MITRE T1557",
          dw.feed(drow("10.0.0.66", ["ACK"])) is True)
    a = find(new_alerts(n), "Unknown DHCP server 10.0.0.66")
    check("...with the client and the address it was given", a is not None and a["severity"] == "critical"
          and a["details"]["client_mac"] == "aa:bb:cc:00:00:99" and "T1557" in tags(a))
    check("...and only once a day per server", dw.feed(drow("10.0.0.66", ["OFFER"])) is False
          and dw.feed(drow("10.0.0.67", ["OFFER"])) is True)
    W.add("dhcp_servers", "10.0.0.68", "selftest")
    dw.trusted._loaded = 0.0
    check("a server added to the watchlist is trusted at once, without restarting the service",
          dw.feed(drow("10.0.0.68", ["OFFER"])) is False)

    rw = l2_watch.RaWatch()
    rrow = lambda src, otype=134: {"proto": "icmp", "id.orig_p": otype, "id.orig_h": src, "id.resp_h": "ff02::1"}
    check("a Router Advertisement from the trusted router is ignored", rw.feed(rrow("fe80::1")) is False)
    check("a Router SOLICITATION (a client looking for a router, type 133) is not an advertisement",
          rw.feed(rrow("fe80::dead", 133)) is False and rw.feed(rrow("::")) is False)
    n = len(feed())
    check("an advertisement from an unknown source is a medium alert", rw.feed(rrow("fe80::dead")) is True)
    a = find(new_alerts(n), "Unknown IPv6 router fe80::dead")
    check("...tagged MITRE T1557", a is not None and a["severity"] == "medium" and "T1557" in tags(a))

    zdir = tmp / "zeek_l2"
    (zdir / "current").mkdir(parents=True)
    now_l2 = time.time()
    (zdir / "current" / "dhcp.log").write_text(
        "#separator \\x09\n#fields\tts\tmac\tserver_addr\tmsg_types\tassigned_addr\n#types\ttime\tstring\taddr\tvector[string]\taddr\n"
        + "".join(f"{now_l2}\taa:aa:aa:aa:aa:aa\t10.9.0.1\tOFFER,ACK\t10.9.0.{i}\n" for i in range(6))
        + f"{now_l2}\tbb:bb:bb:bb:bb:bb\t10.9.0.2\tOFFER\t10.9.0.99\n{now_l2}\tcc:cc:cc:cc:cc:cc\t-\tDISCOVER\t-\n")
    (zdir / "current" / "conn.log").write_text(
        "#separator \\x09\n#fields\tts\tid.orig_h\tid.orig_p\tid.resp_h\tid.resp_p\tproto\n#types\ttime\taddr\tport\taddr\tport\tenum\n"
        f"{now_l2}\tfe80::aa\t134\tff02::1\t0\ticmp\n"                       # a recent advertisement: learned
        f"{now_l2 - 10 * 86400}\tfe80::bb\t134\tff02::1\t0\ticmp\n"           # ten days old: a re-addressed router, not learned
        f"{now_l2}\tfe80::cc\t133\tff02::2\t134\ticmp\n")                     # a solicitation: not a router
    check("DHCP servers are learned from the history, with how often each was seen; a DISCOVER is not a server",
          l2_watch.learn_dhcp_servers(zdir) == {"10.9.0.1": 6, "10.9.0.2": 1}, str(l2_watch.learn_dhcp_servers(zdir)))
    check("IPv6 routers are learned only from recent days and only from advertisements (a router can be re-addressed)",
          l2_watch.learn_ra_sources(zdir) == {"fe80::aa": 1}, str(l2_watch.learn_ra_sources(zdir)))
    base_state = l2_watch.State(tmp / "l2_base_state.json")
    n = len(feed())
    learned = l2_watch.ensure_baselines(base_state, zdir)
    a = find(new_alerts(n), "Baseline established: 2 trusted DHCP server(s)")
    check("first start learns both trusted lists into the watchlists and raises one normal alert that lists them",
          sorted(learned) == ["dhcp_servers", "ra_sources"] and a is not None and a["severity"] == "normal"
          and "10.9.0.2" in a["details"]["rarely_seen"] and "10.9.0.1" in W.get("dhcp_servers")["entries"]
          and "fe80::aa" in W.get("ra_sources")["entries"])
    n = len(feed())
    check("...and never again (an administrator who empties a list is not overruled)",
          l2_watch.ensure_baselines(base_state, zdir) == [] and not new_alerts(n))

    (tmp / "targets_l2.conf").write_text("10.7.0.0/24\n# comment\n192.168.9.0/24\n10.8.0.1/32\n")
    (tmp / "resolv_l2.conf").write_text("nameserver 10.7.0.14\nnameserver ::1\nsearch x\n")
    crit = l2_watch.critical_addresses(tmp / "targets_l2.conf", tmp / "resolv_l2.conf")
    check("critical addresses are each VLAN's .1 gateway, the DNS servers and the trusted DHCP servers (not a /32, not IPv6)",
          crit.get("10.7.0.1") == "gateway" and crit.get("192.168.9.1") == "gateway" and crit.get("10.7.0.14") == "DNS server"
          and crit.get("10.0.0.51") == "DHCP server" and "10.8.0.1" not in crit and "::1" not in crit, str(crit))

    real_crit = l2_watch.critical_addresses
    l2_watch.critical_addresses = lambda *a_, **k_: {"10.7.0.1": "gateway", "10.7.0.2": "gateway"}
    try:
        owners_l2 = tmp / "l2_owners.json"
        scan = lambda gw_mac, extra=None: {"10.7.0.0/24": {
            **{m_: {"ip": "10.7.0.1", "vendor": "GW Inc", "iface": "eth0"} for m_ in ([gw_mac] + (extra or []))},
            "aa:00:00:00:00:50": {"ip": "10.7.0.50", "vendor": "PC", "iface": "eth0"},
            "00:00:5e:00:01:07": {"ip": "10.7.0.2", "vendor": "VRRP", "iface": "eth0"}}}
        n = len(feed())
        check("the first sighting of a critical address is recorded silently",
              R.check_critical_addresses(scan("aa:00:00:00:00:01"), owners_l2) == 0)
        check("the same owner on the next scan is silent", R.check_critical_addresses(scan("aa:00:00:00:00:01"), owners_l2) == 0)
        check("a different MAC on the gateway is a critical alert (ARP poisoning / rogue gateway), tagged T1557.002",
              R.check_critical_addresses(scan("aa:00:00:00:00:02"), owners_l2) == 1)
        a = find(new_alerts(n), "Critical address 10.7.0.1 (gateway) is now answered by a different MAC")
        check("...naming both MACs and saying it may be a legitimate failover",
              a is not None and a["severity"] == "critical" and "aa:00:00:00:00:01" in a["title"] and "failover" in a["description"]
              and "T1557.002" in tags(a))
        n = len(feed())
        check("a VRRP failover (only the physical MAC in parentheses changes) is NOT an alert",
              R.check_critical_addresses({"10.7.0.0/24": {"00:00:5e:00:01:07 (11:22:33:44:55:66)":
                                          {"ip": "10.7.0.2", "vendor": "VRRP", "iface": "eth0"}}}, owners_l2) == 0
              and R.check_critical_addresses({"10.7.0.0/24": {"00:00:5e:00:01:07 (66:55:44:33:22:11)":
                                             {"ip": "10.7.0.2", "vendor": "VRRP", "iface": "eth0"}}}, owners_l2) == 0)
        check("two MACs answering for a critical address in one scan is a critical alert",
              R.check_critical_addresses(scan("aa:00:00:00:00:02", ["aa:00:00:00:00:03"]), owners_l2) == 1
              and find(new_alerts(n), "Two MAC addresses answer for critical address 10.7.0.1") is not None)
    finally:
        l2_watch.critical_addresses = real_crit

    tx = 0x5A5A
    q = PC.build_llmnr_query("qwerty12", tx)
    check("the LLMNR canary is a well-formed query for a type-A name",
          q[:2] == _struct.pack(">H", tx) and q[4:6] == b"\x00\x01" and b"\x08qwerty12\x00" in q and q.endswith(b"\x00\x01\x00\x01"))
    nb = PC.build_nbns_query("wpadtest", tx)
    check("the NBT-NS canary is a 50-byte name query (broadcast flag, type NB)", len(nb) == 50 and nb[2:4] == b"\x01\x10" and nb[-4:-2] == b"\x00\x20")
    resp_hdr = lambda tid, flags, an: _struct.pack(">HHHHHH", tid, flags, 1, an, 0, 0) + b"x" * 4
    check("only a positive answer to OUR question counts (not another id, not a query, not a WINS 'name not found', not empty)",
          PC.parse_llmnr_response(resp_hdr(tx, 0x8000, 1), tx) == {"protocol": "LLMNR", "answers": 1}
          and PC.parse_nbns_response(resp_hdr(tx, 0x8500, 1), tx) == {"protocol": "NBT-NS", "answers": 1}
          and PC.parse_llmnr_response(resp_hdr(tx + 1, 0x8000, 1), tx) is None
          and PC.parse_llmnr_response(resp_hdr(tx, 0x0000, 1), tx) is None
          and PC.parse_nbns_response(resp_hdr(tx, 0x8003, 0), tx) is None
          and PC.parse_llmnr_response(resp_hdr(tx, 0x8000, 0), tx) is None)
    fake = _socket.socket(_socket.AF_INET, _socket.SOCK_DGRAM)
    fake.bind(("127.0.0.1", 0))
    fake_port = fake.getsockname()[1]

    def poisoner():                       # answers whatever it is asked, like Responder
        fake.settimeout(3)
        try:
            data_, addr_ = fake.recvfrom(2048)
            fake.sendto(data_[:2] + _struct.pack(">HHHHH", 0x8000, 1, 1, 0, 0) + b"xxxx", addr_)
        except OSError:
            pass
    th = _threading.Thread(target=poisoner)
    th.start()
    hits_ = PC.probe("127.0.0.1", ("127.0.0.1", fake_port), PC.build_llmnr_query("abc", 0x77), 0x77, PC.parse_llmnr_response,
                     set(), wait=1.0)
    th.join()
    check("the canary detects a poisoner that answers its made-up name (real sockets, fake poisoner)",
          [h_["responder"] for h_ in hits_] == ["127.0.0.1"] and hits_[0]["protocol"] == "LLMNR")
    silent = _socket.socket(_socket.AF_INET, _socket.SOCK_DGRAM)
    silent.bind(("127.0.0.1", 0))
    check("...and finds nothing when nobody answers",
          PC.probe("127.0.0.1", ("127.0.0.1", silent.getsockname()[1]), PC.build_llmnr_query("abc", 1), 1,
                   PC.parse_llmnr_response, set(), wait=0.5) == [])
    silent.close()
    fake.close()
    hit_ = {"responder": "10.7.0.66", "protocol": "LLMNR", "iface": "eth9", "name": "qwerty12", "network": "10.7.0.5/24"}
    n = len(feed())
    check("an answer to the canary is a critical alert tagged MITRE T1557.001 as OBSERVED (direct evidence)",
          PC.report(hit_) is True)
    a = find(new_alerts(n), "LLMNR/NBT-NS poisoner answering on eth9: 10.7.0.66")
    check("...that names the invented name and says what to do",
          a is not None and a["severity"] == "critical" and "qwerty12" in a["description"] and "LLMNR and NBT-NS" in a["description"]
          and any(t_["technique"] == "T1557.001" and t_["basis"] == "observed" for t_ in a["details"]["mitre"]))
    check("...once a day per responder", PC.report(hit_) is False)
    import service_watchdog
    check("an optional unit is watched only once it is installed (no false 'service down' before the installer runs)",
          ("soc-l2-watch" in service_watchdog.SERVICES) == Path("/etc/systemd/system/soc-l2-watch.service").exists())

    late_dhcp, late_conn = tmp / "late_dhcp.log", tmp / "late_conn.log"
    empty_zeek = tmp / "zeek_empty"
    empty_zeek.mkdir()
    l2 = subprocess.Popen([sys.executable, str(KALI / "l2_to_alerts.py"), "--follow", "--dhcp-log", str(late_dhcp),
                           "--conn-log", str(late_conn)], env=dict(os.environ, SOC_ZEEK_DIR=str(empty_zeek)),
                          stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        time.sleep(1.5)
        check("the l2 follower waits for logs that do not exist yet instead of exiting", l2.poll() is None)
        late_dhcp.write_text("#separator \\x09\n#fields\tts\tmac\tserver_addr\tmsg_types\tassigned_addr\n"
                             "#types\ttime\tstring\taddr\tvector[string]\taddr\n"
                             "1.0\taa:aa:aa:aa:aa:aa\t10.0.0.99\tOFFER\t10.0.0.120\n")
        late_conn.write_text("#separator \\x09\n#fields\tts\tid.orig_h\tid.orig_p\tid.resp_h\tid.resp_p\tproto\n"
                             "#types\ttime\taddr\tport\taddr\tport\tenum\n1.0\tfe80::abcd\t134\tff02::1\t0\ticmp\n")
        deadline, got_dhcp, got_ra = time.time() + 8, False, False
        while time.time() < deadline and not (got_dhcp and got_ra):
            got_dhcp = got_dhcp or find(feed(), "Unknown DHCP server 10.0.0.99") is not None
            got_ra = got_ra or find(feed(), "Unknown IPv6 router fe80::abcd") is not None
            time.sleep(0.3)
        check("...then raises both alerts from logs that appeared after it started (read from their first line)",
              got_dhcp and got_ra, f"dhcp: {got_dhcp}, ra: {got_ra}")
        check("...and publishes its parse health", all(reader_health.status(n_) is not None for n_ in ("l2_dhcp", "l2_conn")))
    finally:
        l2.terminate()
        try:
            l2.wait(timeout=5)
        except subprocess.TimeoutExpired:
            l2.kill()

    # ---- (added) the API over HTTPS (soc-api-tls.service): a CA-signed certificate like make_api_cert.sh's --------
    group("api tls")
    import socket
    import ssl
    import urllib.error    # (inner() imports urllib locally further down, so the name is local here too)
    import urllib.request
    tdir = tmp / "tls"
    tdir.mkdir()
    ossl = lambda *a: subprocess.run(["openssl", *a], cwd=tdir, capture_output=True, text=True, timeout=30)  # noqa: E731
    ossl("req", "-x509", "-newkey", "ec", "-pkeyopt", "ec_paramgen_curve:P-256", "-nodes", "-days", "2", "-keyout", "ca.key",
         "-out", "ca.pem", "-subj", "/CN=selftest CA", "-addext", "basicConstraints=critical,CA:TRUE,pathlen:0",
         "-addext", "keyUsage=critical,keyCertSign,cRLSign", "-addext", "subjectKeyIdentifier=hash")
    ossl("req", "-newkey", "ec", "-pkeyopt", "ec_paramgen_curve:P-256", "-nodes", "-keyout", "api.key", "-out", "api.csr",
         "-subj", "/CN=localhost")
    (tdir / "ext").write_text("basicConstraints=critical,CA:FALSE\nkeyUsage=critical,digitalSignature\n"
                              "extendedKeyUsage=serverAuth\nsubjectKeyIdentifier=hash\nauthorityKeyIdentifier=keyid\n"
                              "subjectAltName=IP:127.0.0.1\n")
    ossl("x509", "-req", "-in", "api.csr", "-CA", "ca.pem", "-CAkey", "ca.key", "-CAcreateserial", "-days", "2",
         "-sha256", "-out", "api.pem", "-extfile", "ext")
    with socket.socket() as sk:
        sk.bind(("127.0.0.1", 0))
        tport = sk.getsockname()[1]
    tproc = subprocess.Popen([sys.executable, str(SCRIPTS / "soc_api.py")], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                             env=dict(os.environ, SOC_API_HOST="127.0.0.1", SOC_API_PORT=str(tport),
                                      SOC_API_TLS_CERT=str(tdir / "api.pem"), SOC_API_TLS_KEY=str(tdir / "api.key")))
    try:
        trusted = ssl.create_default_context(cafile=str(tdir / "ca.pem"))   # strict X.509 checks, like the backend
        ok_tls = None
        for _ in range(40):
            try:
                with urllib.request.urlopen(f"https://127.0.0.1:{tport}/api/health", context=trusted, timeout=2) as r:
                    ok_tls = r.status
                break
            except OSError:
                time.sleep(0.25)
        check("the API serves HTTPS with a certificate the backend can verify strictly (CA file, like KALI_API_CA_CERT)",
              ok_tls == 200, str(ok_tls))
        try:
            urllib.request.urlopen(f"https://127.0.0.1:{tport}/api/health", context=ssl.create_default_context(), timeout=3)
            refused = False
        except OSError:
            refused = True
        check("...and a client that does not trust the CA is refused", refused)
        try:
            urllib.request.urlopen(f"http://127.0.0.1:{tport}/api/health", timeout=3)
            plain = True
        except OSError:
            plain = False
        check("...and plain HTTP on the TLS port gets nothing", not plain)
        with socket.create_connection(("127.0.0.1", tport), timeout=3):
            stall_start = time.monotonic()           # a client that connects and never handshakes...
            try:
                with urllib.request.urlopen(f"https://127.0.0.1:{tport}/api/health", context=trusted, timeout=5) as r:
                    served = r.status == 200
            except OSError:
                served = False
            check("...and one stalled client does not block the others (handshake runs per connection)",
                  served and time.monotonic() - stall_start < 5)
        import cert_expiry as CE
        na_ = CE.not_after_file(tdir / "api.pem")
        check("cert_expiry reads a certificate's expiry from the file and from the running HTTPS API",
              na_ is not None and CE.not_after_live(f"127.0.0.1:{tport}", tdir / "ca.pem") == na_
              and CE.not_after_file(tdir / "missing.pem") is None)
    finally:
        tproc.terminate()
        tproc.wait(timeout=10)
    import cert_expiry as CE
    from datetime import timedelta
    now_ = datetime(2026, 9, 24, tzinfo=timezone.utc)
    exp = lambda days: {"api": ("API cert", now_ + timedelta(days=days))}  # noqa: E731
    a1, s1 = CE.evaluate(exp(400), {}, now_)
    a2, s2 = CE.evaluate(exp(25), s1, now_)
    a3, s3 = CE.evaluate(exp(24), s2, now_)
    a4, s4 = CE.evaluate(exp(5), s3, now_)
    a5, _ = CE.evaluate(exp(800), s4, now_)
    check("certificate expiry: nothing far out, medium at 30 days, once, critical at 7, and a note when renewed",
          a1 == [] and [x_["severity"] for x_ in a2] == ["medium"] and a3 == []
          and [x_["severity"] for x_ in a4] == ["critical"] and [x_["severity"] for x_ in a5] == ["normal"]
          and "renewed" in a5[0]["title"], str([a1, a2, a3, a4, a5])[:300])
    a6, _ = CE.evaluate({"api": ("API cert", now_ - timedelta(days=1))}, {}, now_)
    check("...an expired one says so, and an unreadable one changes nothing",
          "EXPIRED" in a6[0]["title"] and a6[0]["severity"] == "critical"
          and CE.evaluate({"api": ("API cert", None)}, {"api": "medium"}, now_) == ([], {"api": "medium"}))

    # ---- (added) API server with read and write keys ------------------------------
    group("api")
    import socket
    import urllib.error
    (tmp / "api_keys.json").write_text(json.dumps({"keys": [{"token": "t-read", "user": "reader", "role": "read"},
                                                             {"token": "t-write", "user": "writer", "role": "write"}]}))
    with socket.socket() as sk:
        sk.bind(("127.0.0.1", 0))
        port = sk.getsockname()[1]
    proc = subprocess.Popen([sys.executable, str(SCRIPTS / "soc_api.py")], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                            env=dict(os.environ, SOC_API_HOST="127.0.0.1", SOC_API_PORT=str(port)))

    def _call(path, token="t-read", method="GET", body=None):
        req = urllib.request.Request(f"http://127.0.0.1:{port}{path}", method=method,
                                     data=json.dumps(body).encode() if body is not None else None,
                                     headers={"Authorization": f"Bearer {token}", "Content-Type": "application/json"} if token else {})
        try:
            with urllib.request.urlopen(req, timeout=15) as r:
                return r.status, json.load(r)
        except urllib.error.HTTPError as e:
            return e.code, None
    try:
        for _ in range(40):
            try:
                urllib.request.urlopen(f"http://127.0.0.1:{port}/api/health", timeout=1)
                break
            except OSError:
                time.sleep(0.25)
        check("the API answers health without a key", _call("/api/health", None)[0] == 200)
        check("the API refuses a request without a key", _call("/api/alerts", None)[0] == 401)
        code, d = _call("/api/alerts?limit=5")
        check("the alerts endpoint serves the feed", code == 200 and d["count"] >= 1)
        sync_id = d["alerts"][0]["id"]
        mark = datetime.now(timezone.utc).isoformat()
        code, d = _call(f"/api/alerts/{sync_id}/status", "t-write", "POST", {"status": "acknowledged", "note": "looking"})
        code, d = _call("/api/alerts?status_since=" + urllib.parse.quote(mark))
        got_ = next((x_ for x_ in (d or {}).get("alerts", []) if x_["id"] == sync_id), None)
        check("status_since returns an alert whose status just changed, with who, when and why",
              code == 200 and got_ is not None and got_["status"] == "acknowledged" and got_["status_actor"] == "writer"
              and got_["status_note"] == "looking" and got_.get("status_updated", "") >= mark, str(got_))
        # AD views (the inventory written by the "ad inventory" group above: OLDPC and jdoe, a Domain Admin)
        code, d = _call("/api/ad/summary")
        check("GET /api/ad/summary serves counts by status, the privileged members and the policy",
              code == 200 and d["computers"]["total"] == 1 and d["privileged"]["Domain Admins"][0]["sam"] == "jdoe",
              str(d)[:300])
        code, d = _call("/api/ad/users?privileged=1")
        check("GET /api/ad/users?privileged=1 lists the admins with their groups",
              code == 200 and [u["sam"] for u in d["users"]] == ["jdoe"] and d["users"][0]["privileged_groups"])
        code, d = _call("/api/ad/computers?q=oldpc")
        check("GET /api/ad/computers finds a computer, with its status and Windows support",
              code == 200 and d["count"] == 1 and d["computers"][0]["status"] in ("active", "stale")
              and d["computers"][0]["support"]["status"] == "unsupported", str(d)[:300])
        check("an unknown AD status is a 400, and GET /api/ad/history answers",
              _call("/api/ad/computers?status=bogus")[0] == 400 and _call("/api/ad/history")[0] == 200)
        code, d = _call("/api/alerts?status_since=2999-01-01T00:00:00Z")
        check("...and nothing for changes that have not happened", code == 200 and d["count"] == 0)
        _call(f"/api/alerts/{sync_id}/status", "t-write", "POST", {"status": "open"})
        code, d = _call("/api/alerts?status_since=" + urllib.parse.quote(mark))
        got_ = next((x_ for x_ in (d or {}).get("alerts", []) if x_["id"] == sync_id), {})
        check("a later change without a note clears the old note", got_.get("status") == "open" and "status_note" not in got_,
              str(got_))
        check("incidents, entities, metrics, sources, detections and playbooks answer",
              all(_call(p)[0] == 200 for p in ("/api/incidents", "/api/entities", "/api/metrics", "/api/sources",
                                               "/api/detections", "/api/playbooks", "/api/suppressions")))
        code, d = _call("/api/self")
        check("GET /api/self returns the hostname and this box's addresses",
              code == 200 and bool(d.get("hostname")) and isinstance(d.get("addresses"), list)
              and all({"ip", "interface"} <= set(x) for x in d["addresses"]))
        check("a search through the API works", _call("/api/search?source=alerts&q=batch&since=1h")[0] == 200)
        check("a path-traversal search source is rejected", _call("/api/search?source=zeek:../../etc/passwd&q=x")[0] == 400)
        check("a read-only key cannot change an incident", _call(f"/api/incidents/{inc_id}", "t-read", "POST", {"comment": "x"})[0] == 403)
        code, d = _call(f"/api/incidents/{inc_id}", "t-write", "POST", {"comment": "through the API"})
        check("a write key can comment, attributed to its user", code == 200 and d["comments"][-1]["by"] == "writer")
        code, d = _call(f"/api/incidents/{inc_id}/graph")
        check("GET /api/incidents/<id>/graph serves a node/edge graph",
              code == 200 and any(n["kind"] == "incident" for n in d["nodes"]) and d["edges"])
        check("the incident graph 404s for an unknown incident", _call("/api/incidents/no-such-incident/graph")[0] == 404)
        code, d = _call("/api/entities?limit=1")
        ekey = d["entities"][0]["key"] if code == 200 and d["entities"] else None
        if ekey:
            code, d = _call(f"/api/entities/{ekey}/graph")
            check("GET /api/entities/<ref>/graph serves a node/edge graph",
                  code == 200 and any(n["id"] == ekey for n in d["nodes"]))
        check("the entity graph 404s for an unknown entity", _call("/api/entities/ip:198.51.100.251/graph")[0] == 404)
        check("a read-only key cannot create a suppression", _call("/api/suppressions", "t-read", "POST", {"reason": "x"})[0] == 403)

        code, d = _call("/api/watchlists")
        check("GET /api/watchlists lists all nine", code == 200 and {w["name"] for w in d["watchlists"]} ==
              {"trusted_ips", "bad_ips", "bad_domains", "bad_hashes", "dhcp_servers", "ra_sources", "sensitive_vlans",
               "on_leave_accounts",
               "untrusted_vlans"})
        code, d = _call("/api/watchlists/trusted_ips")
        check("GET /api/watchlists/<name> serves one list", code == 200 and d["name"] == "trusted_ips")
        check("GET on an unknown watchlist name is a 404", _call("/api/watchlists/does-not-exist")[0] == 404)
        check("a read-only key cannot add to a watchlist",
              _call("/api/watchlists/trusted_ips", "t-read", "POST", {"entry": "198.51.100.9"})[0] == 403)
        code, d = _call("/api/watchlists/trusted_ips", "t-write", "POST", {"entry": "198.51.100.9"})
        check("a write key can add an entry", code == 201 and "198.51.100.9" in d["entries"])
        code, d = _call("/api/watchlists/trusted_ips?entry=198.51.100.9", "t-write", "DELETE")
        check("a write key can remove an entry", code == 200 and "198.51.100.9" not in d["updated"]["entries"])

        code, d = _call("/api/activity?limit=500")
        check("GET /api/activity serves a merged feed with several categories present",
              code == 200 and d["count"] > 0 and {"watchlist", "suppression"} <= {r["category"] for r in d["activity"]})
        check("an unknown category is a 400", _call("/api/activity?category=not-a-real-category")[0] == 400)
        code, d = _call("/api/activity?category=watchlist")
        check("the category filter actually filters", code == 200 and d["activity"]
              and all(r["category"] == "watchlist" for r in d["activity"]))

        code, d = _call("/api/mitre")
        check("GET /api/mitre serves the coverage matrix", code == 200 and d["summary"]["techniques_tracked"] > 0
              and {t["tactic"] for t in d["tactics"]} >= {"Credential Access", "Lateral Movement"})
        code, d = _call("/api/reports/weekly")
        check("GET /api/reports/weekly builds a live report", code == 200 and {"alerts", "exposure", "coverage", "attention"} <= set(d))
        code, d = _call("/api/reports/weekly?format=markdown")
        check("...and the same as Markdown on request", code == 200 and d["markdown"].startswith("# Sentinel SOC weekly report"))
        check("a saved report that does not exist is a 404, not a crash", _call("/api/reports/weekly?date=1999-01-01")[0] == 404)
        check("a malformed or path-like date is refused", _call("/api/reports/weekly?date=../../etc/passwd")[0] == 404)
        code, d = _call("/api/reports")
        check("GET /api/reports lists the saved reports", code == 200 and isinstance(d["reports"], list))
    finally:
        proc.terminate()
        try:
            proc.wait(timeout=5)
        except subprocess.TimeoutExpired:
            proc.kill()

    # ---- (added) unified activity feed (soc_activity.py) ----------------------
    group("activity")
    import soc_activity
    check("the 7 audit-log sources are registered",
          set(soc_activity.CATEGORIES) == {"alert_status", "asset_annotation", "incident",
                                           "playbook_run", "suppression", "watchlist", "alert_aging"})
    rows = soc_activity.feed(limit=500)
    check("the feed merges entries from multiple sources, most-recent-first",
          len(rows) > 0 and all(rows[i]["at"] >= rows[i + 1]["at"] for i in range(len(rows) - 1)))
    check("every row has the normalized shape", all({"at", "actor", "category", "action", "summary", "ref"} <= set(r) for r in rows))
    by_actor = soc_activity.feed(actor="selftest", limit=500)
    check("the actor filter works", by_actor and all(r["actor"] == "selftest" for r in by_actor))
    try:
        soc_activity.feed(since="not-a-duration")
        check("a malformed since is rejected", False)
    except ValueError:
        check("a malformed since is rejected", True)
    # a deliberately ancient entry, appended directly (bypassing every writer above), proves
    # `since` actually excludes old rows instead of just accepting any value silently
    old_path = soc_core.DATA_DIR / "watchlist_log.jsonl"
    with old_path.open("a") as fh:
        fh.write(json.dumps({"at": "2000-01-01T00:00:00+00:00", "actor": "selftest-ancient",
                             "action": "added", "watchlist": "trusted_ips", "entry": "0.0.0.0/32"}) + "\n")
    check("`since` excludes an entry older than the window",
          "selftest-ancient" not in {r["actor"] for r in soc_activity.feed(since="1h", limit=500)}
          and "selftest-ancient" in {r["actor"] for r in soc_activity.feed(limit=500)})

    # ---- (added) alert geolocation: two databases, country from the newer one ------------------------
    group("geolocation")
    import geoip_enrich as GE
    from types import SimpleNamespace as NS

    class FakeCountry:
        def __init__(self, table): self.table = table
        def country(self, ip):
            if ip not in self.table: raise KeyError(ip)
            return NS(country=NS(iso_code=self.table[ip][0], name=self.table[ip][1]))

    class FakeCity:
        def __init__(self, table): self.table = table
        def city(self, ip):
            if ip not in self.table: raise KeyError(ip)
            cc, nm, city, lat, lon = self.table[ip]
            return NS(country=NS(iso_code=cc, name=nm), city=NS(name=city), location=NS(latitude=lat, longitude=lon))
    saved_readers = dict(GE._readers)
    GE._next_try.update(city=1e18, country=1e18)          # do not go looking for the real database files
    try:
        GE._readers["country"] = FakeCountry({"5.5.5.5": ("DE", "Germany"), "5.5.5.6": ("NL", "Netherlands"),
                                              "5.5.5.7": ("RU", "Russia")})
        GE._readers["city"] = FakeCity({"5.5.5.5": ("DE", "Germany", "Berlin", 52.5, 13.4),
                                        "5.5.5.6": ("HK", "Hong Kong", "Kowloon", 22.3, 114.2),
                                        "5.5.5.8": ("US", "United States", None, 37.0, -95.0)})
        g = GE.geolocate("5.5.5.5")
        check("country from the country database, city and coordinates added when the city database agrees",
              g == {"country_code": "DE", "country": "Germany", "city": "Berlin", "lat": 52.5, "lon": 13.4, "precision": "city"}, str(g))
        g = GE.geolocate("5.5.5.6")
        check("when the old city database disagrees on the country, its city and coordinates are dropped (never contradicts)",
              g == {"country_code": "NL", "country": "Netherlands", "precision": "country"}, str(g))
        check("a block only the newer country database knows still gets a country",
              GE.geolocate("5.5.5.7") == {"country_code": "RU", "country": "Russia", "precision": "country"})
        g = GE.geolocate("5.5.5.8")
        check("an address only the city database knows falls back to it, and no city name means precision 'country'",
              g and g["country_code"] == "US" and g["precision"] == "country" and "city" not in g, str(g))
        check("private, loopback and unknown addresses give nothing",
              GE.geolocate("192.168.1.1") is None and GE.geolocate("127.0.0.1") is None and GE.geolocate("5.5.5.99") is None
              and GE.geolocate(None) is None)
        n = len(feed())
        soc_core.emit_alert(soc_core.Alert(type="intrusion", severity="normal", title="Geo probe (selftest)", detector="selftest",
                                           source_ip="5.5.5.5", description="x"), echo=False)
        a = find(new_alerts(n), "Geo probe (selftest)")
        check("a real alert from a public IP carries details.geo through emit_alert",
              a is not None and a["details"].get("geo", {}).get("country_code") == "DE")
        n = len(feed())
        soc_core.emit_alert(soc_core.Alert(type="intrusion", severity="normal", title="Geo probe private (selftest)",
                                           detector="selftest", source_ip="10.1.2.3", description="x"), echo=False)
        check("...and an internal IP carries none", "geo" not in find(new_alerts(n), "Geo probe private (selftest)")["details"])
        GE._readers.update(country=None, city=None)
        check("with no database at all it is a quiet no-op (and the doctor's status says so)",
              GE.geolocate("5.5.5.5") is None and GE.status()["available"] is False)

        # the health check: geolocation must never again be off without anyone being told
        import source_health as SH
        import time as _t
        prob = SH.geoip_problem([])
        check("source health reports missing databases, naming the files", prob and "MISSING" in prob and "GeoLite2" in prob, str(prob))
        GE._readers["country"] = FakeCountry({"8.8.8.8": ("US", "United States")})
        iso = lambda ago_h: datetime.fromtimestamp(_t.time() - ago_h * 3600, tz=timezone.utc).isoformat()
        mk = lambda **kw: {"title": "x", "source_ip": "5.5.5.5", "timestamp": iso(1), "details": {}, **kw}
        check("working databases and no recent public-IP alerts: healthy", SH.geoip_problem([]) is None)
        check("a recent public-IP alert WITHOUT a location is a problem (a service on old code, or a broken step)",
              "carry no location" in (SH.geoip_problem([mk()]) or ""))
        check("...not when it has one, is a test, is internal, or is older than 6 hours",
              SH.geoip_problem([mk(details={"geo": {"country_code": "DE"}}), mk(test=True), mk(source_ip="10.0.0.5"),
                                mk(timestamp=iso(7))]) is None)
        GE._readers["country"] = FakeCountry({"8.8.8.8": ("DE", "Germany")})
        check("a lookup that returns the wrong country for the probe address is a problem too",
              SH.geoip_problem([]) is not None)
        check("the source is registered and shows as a source with a reason when unhealthy",
              any(x["id"] == "geolocation" and x["kind"] == "geoip" for x in SH.SOURCES)
              and next(r for r in SH.check() if r["id"] == "geolocation")["status"] == "stale")
    finally:
        GE._readers.update(saved_readers)
        GE._next_try.update(city=0.0, country=0.0)

    # ---- (added) identity coverage: IP -> AD computer through DHCP names and the domain's DNS ---------------
    group("identity coverage")
    import ip_names as IPN
    import ad_identity as AID
    from datetime import datetime as _d2, timedelta as _t2, timezone as _z2
    n_now = _d2.now(_z2.utc)
    ago = lambda days: (n_now - _t2(days=days)).isoformat()
    inv_ = {"computers": {
        "PC-ONE": {"name": "PC-ONE", "dns": "pc-one.era.local", "enabled": True, "last_logon": ago(3), "ou": "Calgary/A"},
        "PC-TWO": {"name": "PC-TWO", "dns": "pc-two.era.local", "enabled": True, "last_logon": ago(3)},
        "PC-DHCP": {"name": "PC-DHCP", "dns": "pc-dhcp.era.local", "enabled": True, "last_logon": ago(3)},
        "PC-OLD": {"name": "PC-OLD", "dns": "pc-old.era.local", "enabled": True, "last_logon": ago(200)},
        "PC-OFF": {"name": "PC-OFF", "dns": "pc-off.era.local", "enabled": False, "last_logon": ago(3)},
        "SHARED-A": {"name": "SHARED-A", "dns": "shared-a.era.local", "enabled": True, "last_logon": ago(3)},
        "SHARED-B": {"name": "SHARED-B", "dns": "shared-b.era.local", "enabled": True, "last_logon": ago(3)},
        "CONFLICT": {"name": "CONFLICT", "dns": "conflict.era.local", "enabled": True, "last_logon": ago(3)},
        "MULTI": {"name": "MULTI", "dns": "multi.era.local", "enabled": True, "last_logon": ago(3)},
        "WEIRD": {"name": "WEIRD", "dns": "weird.era.local", "enabled": True, "last_logon": ago(3)}}}
    asset = lambda ip, **kw: {"ip": ip, "last_seen": ago(0.1), **kw}
    assets_ = {"a1": asset("10.9.0.1"),
               "a2": asset("10.9.0.3", dhcp_hostname="pc-dhcp"),
               "a3": asset("10.9.0.9", scan_facts={"windows": {"NetBIOS_Computer_Name": "HP-PRINTER"}}),
               "a4": asset("10.9.0.20", last_seen=ago(0.1)),
               "a5": asset("10.9.0.21", last_seen=ago(30))}
    dns_ = {"pc-one.era.local": ["10.9.0.1"], "pc-two.era.local": ["10.9.0.2"], "pc-dhcp.era.local": ["10.9.0.30"],
            "pc-old.era.local": ["10.9.0.4"], "pc-off.era.local": ["10.9.0.5"],
            "shared-a.era.local": ["10.9.0.6"], "shared-b.era.local": ["10.9.0.6"],
            "conflict.era.local": ["10.9.0.9"], "multi.era.local": ["10.9.0.20", "10.9.0.21"],
            "weird.era.local": ["0.0.0.0", "169.254.1.1", "127.0.0.1"]}
    looked = []
    def fake_lookup(fq):
        looked.append(fq)
        if fq == "pc-two.era.local":
            raise OSError("boom")
        return dns_.get(fq, [])
    r_ = IPN.build(inv_, assets_, fake_lookup, n_now)
    m_ = r_["ips"]
    check("an AD computer whose DNS name resolves to one address is mapped, marked via 'dns'",
          m_.get("10.9.0.1") == {"name": "PC-ONE", "via": "dns"}, str(m_.get("10.9.0.1")))
    check("a host name the DHCP server saw is mapped too, marked via 'dhcp', and wins over a DNS record for it",
          m_.get("10.9.0.3") == {"name": "PC-DHCP", "via": "dhcp"} and "10.9.0.30" in m_ and m_["10.9.0.30"]["via"] == "dns")
    check("a computer that is disabled or has not signed in for 90 days is not even looked up",
          "pc-old.era.local" not in looked and "pc-off.era.local" not in looked and "10.9.0.4" not in m_ and "10.9.0.5" not in m_)
    check("an address that two AD computers resolve to is not attributed to either (a stale record must not name the wrong PC)",
          "10.9.0.6" not in m_ and r_["stats"]["skipped"].get("address shared by several AD computers") == 2)
    check("an address the scan or DHCP named as something else is not overridden by DNS",
          "10.9.0.9" not in m_ and r_["stats"]["skipped"].get("the scan/DHCP named a different machine there") == 1)
    check("a multi-homed machine keeps only the address seen in the last 2 days",
          m_.get("10.9.0.20", {}).get("name") == "MULTI" and "10.9.0.21" not in m_)
    check("a lookup that fails, loopback, link-local and unspecified answers are ignored without breaking the pass",
          "10.9.0.2" not in m_ and not any(k.startswith(("0.", "127.", "169.254.")) for k in m_))
    check("the stats add up", r_["stats"]["mapped"] == len(m_) == r_["stats"]["via_dhcp"] + r_["stats"]["via_dns"])

    # the file, the hourly refresh, and how ad_identity uses it
    saved_files = {n: (tmp / n).read_bytes() for n in ("ad_inventory.json", "ip_names.json") if (tmp / n).exists()}
    try:
        (tmp / "ip_names.json").unlink(missing_ok=True)
        (tmp / "ad_inventory.json").unlink(missing_ok=True)
        check("with no AD inventory yet, refresh does nothing and writes nothing",
              IPN.refresh(force=True) is None and not (tmp / "ip_names.json").exists())
        (tmp / "ad_inventory.json").write_text(json.dumps({**inv_, "generated": ago(0.01), "users": {}, "privileged": {}}))
        real_lookup = IPN._dns_lookup
        IPN._dns_lookup = lambda fq: dns_.get(fq, [])
        try:
            IPN.build.__defaults__ = (IPN._dns_lookup, None)          # refresh() uses build()'s default resolver
            first = IPN.refresh(force=True)
        finally:
            IPN._dns_lookup = real_lookup
            IPN.build.__defaults__ = (real_lookup, None)
        check("refresh writes data/ip_names.json, and does not rebuild again within the hour",
              first and (tmp / "ip_names.json").exists() and IPN.refresh() is None
              and json.loads((tmp / "ip_names.json").read_text())["ips"].get("10.9.0.1", {}).get("name") == "PC-ONE")
        (tmp / "ip_names.json").write_text(json.dumps(r_))          # the map built above from the synthetic assets
        rec = lambda ip: {"detector": "kali_scan", "details": {"entities": [{"type": "ip", "value": ip}]}}
        got = AID.identity(rec("10.9.0.1"))
        check("ad_identity names a DNS-matched address, marks it via 'dns', and gives the OU",
              got and got["computers"][0]["name"] == "PC-ONE" and got["computers"][0].get("via") == "dns"
              and got["computers"][0]["ou"] == "Calgary/A", str(got))
        got = AID.identity(rec("10.9.0.3"))
        check("a DHCP-matched address is named without a 'via' mark (it was observed, not inferred)",
              got and got["computers"][0]["name"] == "PC-DHCP" and "via" not in got["computers"][0], str(got))
        check("an address nobody maps still carries no identity", AID.identity(rec("10.9.0.77")) is None)
        # a stale map naming a computer that has since left AD must not read as 'a Windows machine outside the domain'
        (tmp / "ip_names.json").write_text(json.dumps({"ips": {"10.9.0.88": {"name": "GONE-PC", "via": "dns"}}}))
        check("a DNS/DHCP-only name with no AD match adds no computer (only the scan's own Windows names can say "
              "'not in the domain')", AID.identity(rec("10.9.0.88")) is None)
    finally:
        for n in ("ad_inventory.json", "ip_names.json"):
            (tmp / n).unlink(missing_ok=True)
            if n in saved_files:
                (tmp / n).write_bytes(saved_files[n])

    # ---- (added) ATT&CK coverage matrix and the weekly executive report ------------------------------
    group("mitre matrix")
    import mitre_matrix as MM
    import mitre_tags as MT
    m = MM.matrix()
    sm = m["summary"]
    check("covered + limited + gap add up to the techniques tracked",
          sm["covered"] + sm["limited"] + sm["gap"] == sm["techniques_tracked"] > 0)
    check("every technique a detector can tag is on the matrix, under a tactic the matrix knows",
          set(MT.coverage_map()) <= {t["technique"] for tac in m["tactics"] for t in tac["techniques"]}
          and {x for _n, tacs in MT.TECHNIQUES.values() for x in tacs} <= set(MM.TACTICS))
    check("every 'limited' entry and every gap names a real missing data source",
          set(MM.DATA_SOURCES) >= {k for v in MM.LIMITED_BY.values() for k in v}
          and set(MM.DATA_SOURCES) >= {k for g in MM.GAPS.values() for k in g[2]}
          and set(MM.LIMITED_BY) <= set(MT.TECHNIQUES))
    gaps = [t for tac in m["tactics"] for t in tac["techniques"] if t["status"] == "gap"]
    check("a gap says what data it needs and what it would detect", gaps and all(g["needs"] and g["would_detect"] for g in gaps))
    check("data sources are ranked by how many techniques they would improve",
          all(m["data_sources"][i]["techniques_affected"] >= m["data_sources"][i + 1]["techniques_affected"]
              for i in range(len(m["data_sources"]) - 1)))
    check("techniques that raised alerts in the last 30 days are counted (the earlier tests tagged some)",
          sm["with_alerts_30d"] > 0)
    check("an email-only technique is 'limited' by the missing mail source, not 'covered'",
          next(t for tac in m["tactics"] for t in tac["techniques"] if t["technique"] == "T1566")["limited_by"] == ["m365"])
    real_cov, real_tech = MT.coverage_map, dict(MT.TECHNIQUES)
    MT.TECHNIQUES["T1110.003"] = ("Password Spraying", ["Credential Access"])
    MT.coverage_map = lambda: {**real_cov(), "T1110.003": ["login_monitor"]}
    try:
        now_covered = next(t for tac in MM.matrix()["tactics"] for t in tac["techniques"] if t["technique"] == "T1110.003")
        check("a gap becomes covered the moment a detector is written for it (no second list to update)",
              now_covered["status"] != "gap")
    finally:
        MT.coverage_map = real_cov
        MT.TECHNIQUES.clear()
        MT.TECHNIQUES.update(real_tech)

    group("weekly report")
    import weekly_report as WR
    t_now = time.time()

    def _hist(title, sev, det, age_days, extra=None):
        return json.dumps({"timestamp": datetime.fromtimestamp(t_now - age_days * 86400, tz=timezone.utc).isoformat(),
                           "type": "intrusion", "severity": sev, "title": title, "detector": det, "source_ip": None,
                           "details": {}, **(extra or {})})
    with soc_core.ALERTS_LOG.open("a") as fh:
        fh.write("\n".join([_hist(f"File integrity: /srv/wr{i} changed", "critical", "aide", 0.5) for i in range(2500)]
                            + [_hist("Old thing", "medium", "kali_scan", 9),
                               _hist("Demo alert that must not count", "critical", "aide", 0.2, {"test": True}),
                               _hist("RESOLVED: something", "normal", "aide", 0.2)]) + "\n")
    rep = WR.build(t_now)
    check("the report counts real alerts only (not tests, not RESOLVED notices)",
          rep["alerts"]["new"]["now"] >= 2500 and rep["alerts"]["new"]["now"] < 2500 + 2000
          and not any("Demo alert" in c["title"] for c in rep["alerts"]["critical"]))
    check("alerts from the week before are counted for the comparison", rep["alerts"]["new"]["before"] >= 1)
    check("severity counts add up to the total", sum(v["now"] for v in rep["alerts"]["by_severity"].values()) == rep["alerts"]["new"]["now"])
    check("a detector that produced most of the week's alerts in one hour is called out as a single event",
          rep["alerts"]["dominated_by"] and rep["alerts"]["dominated_by"]["detector"] == "aide"
          and rep["alerts"]["dominated_by"]["single_burst"] and any("one event" in x for x in rep["attention"]))
    agg = [c for c in rep["alerts"]["critical"] if c["title"].startswith("File integrity: ")]
    check("2,500 alerts from one event are one line in the critical list, not 2,500 or fifteen",
          len(agg) == 1 and "similar alerts" in agg[0]["title"] and rep["alerts"]["critical_total"] >= 2500)
    md = WR.to_markdown(rep)
    check("the Markdown has the sections a manager reads", all(h in md for h in (
        "## Needs attention", "## The week in numbers", "## Critical alerts this week", "## Known weaknesses still open",
        "## Monitoring health", "## Detection coverage (MITRE ATT&CK)")) and "Calgary time" in md)
    import re as _re
    check("...with thousands separators and a plain-language change ('up N% from M')",
          _re.search(r"\d,\d{3}", md) is not None and _re.search(r"up \d+% from", md) is not None)
    empty = WR.build(t_now + 400 * 86400)
    check("a period with no alerts still produces a report", empty["alerts"]["new"]["now"] == 0 and WR.to_markdown(empty))
    saved = WR.save(rep)
    check("save() writes the Markdown and the JSON, and the JSON reads back",
          saved.exists() and saved.with_suffix(".json").exists()
          and WR.load_saved(saved.stem.replace("weekly_", ""))["period"] == rep["period"])
    check("the saved report is listed", saved.stem.replace("weekly_", "") in [r["date"] for r in WR.list_reports()])
    check("load_saved refuses anything that is not a plain date", WR.load_saved("../../etc/passwd") is None and WR.load_saved("x") is None)

    # ---- (added) reader health: a log the followers cannot parse must not go unnoticed -------------
    group("reader health")
    import reader_health
    import source_health
    from zeek_tsv import ZeekTSVReader
    rd = ZeekTSVReader()
    for _ in range(30):
        rd.feed("#close\t2026-09-23")
        rd.feed("")
    check("comments and blank lines never count as unparsable", rd.unparsable == 0 and rd.unparsable_streak == 0)
    rd.feed("1\t2")
    check("a TSV row that arrives before any header counts as unparsable", rd.unparsable == 1 and rd.unparsable_streak == 1)
    rd.feed('{"ts":1.0,"note":"X"}')
    check("one good line resets the streak", rd.unparsable_streak == 0 and rd.parsed == 1)
    rep = reader_health.Reporter("zeek_to_alerts", rd)
    for _ in range(12):
        rd.feed('{"broken')
        rep.tick()
    real_sources = source_health.SOURCES
    source_health.SOURCES = [s_ for s_ in real_sources if s_["id"] == "zeek_notice_parsing"]
    try:
        st = source_health.check()[0]
        check("a reader whose input stopped parsing is reported unhealthy",
              st["status"] == "stale" and st["unparsable_in_a_row"] >= reader_health.UNHEALTHY_STREAK, str(st))
        n = len(feed()); source_health.run()
        a = find(new_alerts(n), "Log reader cannot parse its input")
        check("...and raises one alert that names the likely cause (regression: Zeek's format flip went unnoticed for two days)",
              a is not None and a["severity"] == "medium" and "TSV" in a["description"])
        n = len(feed()); source_health.run()
        check("...only once while it stays broken", find(new_alerts(n), "Log reader cannot parse") is None)
        rd.feed('{"ts":2.0,"note":"Y"}')
        rep.tick()
        check("a good line ends the streak: healthy again", source_health.check()[0]["status"] == "healthy")
        n = len(feed()); source_health.run()
        check("...with a recovered alert", find(new_alerts(n), "Data source recovered: Zeek notice forwarder") is not None)
    finally:
        source_health.SOURCES = real_sources

    # ---- (added) code freshness: a service still running code older than what is on disk ---------
    group("code freshness")
    import code_freshness as CF
    cf = tmp / "cf"
    (cf / "scripts").mkdir(parents=True)
    (cf / "kali").mkdir()
    (cf / "scripts" / "deep.py").write_text("X = 1\n")
    (cf / "scripts" / "mid.py").write_text("def f():\n    import deep\n")   # a lazy import inside a function still loads
    (cf / "kali" / "svc.py").write_text("import json\nfrom mid import f\n")
    for name_, mtime_ in (("scripts/deep.py", 1000), ("scripts/mid.py", 500), ("kali/svc.py", 400)):
        os.utime(cf / name_, (mtime_, mtime_))
    newest_m, newest_f = CF.newest_code(cf / "kali" / "svc.py", [cf / "scripts", cf / "kali"])
    check("the newest file a service's code is made of is found through nested and lazy imports",
          newest_f.name == "deep.py" and newest_m == 1000, f"{newest_f} {newest_m}")
    check("stale means edited after the service started, and not in the last hour",
          CF.is_stale(100, 5000, 100000) and not CF.is_stale(100, 50, 100000)
          and not CF.is_stale(100, 99999, 100000) and not CF.is_stale(100, 101, 100000))
    real_sysctl = CF._systemctl
    CF._systemctl = lambda unit, *props: ("{ path=/usr/bin/python3 ; argv[]=/usr/bin/python3 "
                                          f"{SUITE}/kali/zeek_to_alerts.py --follow ; ignore_errors=no }}")
    check("the script a unit runs is read from its ExecStart", CF.service_script("x") == SUITE / "kali" / "zeek_to_alerts.py")
    CF._systemctl = lambda unit, *props: "{ path=/usr/bin/suricata ; argv[]=/usr/bin/suricata -c /etc/suricata.yaml ; ignore_errors=no }"
    check("a unit that runs something other than this project's code is skipped", CF.service_script("x") is None)
    CF._systemctl = real_sysctl
    fake = [{"unit": "soc-fake", "started": 1.0, "edited": 2.0, "file": "kali/x.py"}]
    real_stale = CF.stale_services
    CF.stale_services = lambda units, grace=CF.GRACE_SECONDS: fake
    try:
        n = len(feed()); CF.run(["soc-fake"])
        a = find(new_alerts(n), "running older code")
        check("a stale service raises one normal-severity alert with the restart command",
              a is not None and a["severity"] == "normal" and "systemctl restart soc-fake" in a["description"])
        n = len(feed()); CF.run(["soc-fake"])
        check("...and does not repeat while nothing changes", find(new_alerts(n), "running older code") is None)
        fake = []
        CF.run(["soc-fake"])
        check("...and resolves itself once every service runs the current code",
              {x_["id"]: x_ for x_ in feed()}[a["id"]]["status"] == "resolved")
    finally:
        CF.stale_services = real_stale

    group("deploy drift")
    import deploy_drift as DD
    dd_repo, dd_live = tmp / "dd_repo", tmp / "dd_live"
    dd_repo.mkdir()
    dd_live.mkdir()
    (dd_repo / "soc-x.service").write_text("[Service]\n# the topic lives in a drop-in\nExecStart=/bin/true\n")
    (dd_live / "soc-x.service").write_text("[Service]\nEnvironment=NTFY_TOPIC=abc123\n  ExecStart=/bin/true  \n\n")
    (dd_live / "soc-x.service.d").mkdir()
    (dd_live / "soc-x.service.d" / "ntfy.conf").write_text("[Service]\nEnvironment=NTFY_TOPIC=abc123\n")
    check("an installed unit that matches deploy/ apart from the ntfy topic, comments and whitespace is not drift",
          DD.drift(dd_repo, dd_live, check_crontab=False) == [])
    (dd_live / "soc-new.timer").write_text("[Timer]\nOnCalendar=daily\n")
    (dd_live / "soc-x.service").write_text("[Service]\nExecStart=/bin/false\n")
    (dd_repo / "soc-old.timer").write_text("[Timer]\nOnCalendar=daily\n")
    got = DD.drift(dd_repo, dd_live, check_crontab=False)
    check("a unit installed only in /etc, one edited there, and one never installed are each reported",
          sorted(got) == [("changed", "soc-x.service"), ("missing", "soc-new.timer"), ("not installed", "soc-old.timer")], str(got))
    (dd_repo / "soc-old.timer").write_text("[Timer]\nOnCalendar=daily\nEnvironment=NTFY_TOPIC=abc123\n")
    check("a topic written into deploy/ is reported as a secret",
          ("secret", "soc-old.timer") in DD.drift(dd_repo, dd_live, check_crontab=False))
    (dd_repo / "crontab.txt").write_text("17 3 * * * backup\n")
    check("a crontab that differs from deploy/crontab.txt is reported, an equal one is not",
          ("crontab", "crontab.txt") in DD.drift(dd_repo, dd_live, crontab="17 3 * * * other\n")
          and ("crontab", "crontab.txt") not in DD.drift(dd_repo, dd_live, crontab="# m h\n17 3 * * * backup\n"))

    # outdated vs changed: an installed unit equal to an EARLIER committed version is just behind deploy/
    gr = tmp / "dd_git"
    (gr / "deploy").mkdir(parents=True)
    gl = tmp / "dd_git_live"
    gl.mkdir()
    g = lambda *a: subprocess.run(["git", "-C", str(gr), *a], capture_output=True, text=True, timeout=20)  # noqa: E731
    g("init", "-q")
    g("config", "user.email", "selftest@localhost")
    g("config", "user.name", "selftest")
    unit = gr / "deploy" / "soc-y.service"
    unit.write_text("[Service]\nExecStart=/bin/true\n")
    g("add", "-A")
    g("commit", "-qm", "v1")
    unit.write_text("[Service]\nExecStart=/bin/true\nEnvironment=PYTHONUNBUFFERED=1\n")
    g("commit", "-qam", "v2")
    (gl / "soc-y.service").write_text("[Service]\n# old comment\nExecStart=/bin/true\n")
    check("a unit installed from an earlier committed version of deploy/ is 'outdated' (safe to install), not 'changed'",
          DD.drift(gr / "deploy", gl, check_crontab=False) == [("outdated", "soc-y.service")],
          str(DD.drift(gr / "deploy", gl, check_crontab=False)))
    (gl / "soc-y.service").write_text("[Service]\nExecStart=/bin/false\n")
    check("...while a version that was never committed is still 'changed' (an edit made on the machine)",
          DD.drift(gr / "deploy", gl, check_crontab=False) == [("changed", "soc-y.service")])

    # sensor configuration (deploy/sensors/<absolute path>), e.g. a package update replacing Zeek's local.zeek
    sroot = tmp / "dd_sroot"
    zk = gr / "deploy" / "sensors" / "opt" / "zeek" / "site"
    zk.mkdir(parents=True)
    (zk / "local.zeek").write_text("@load base/frameworks\n@load policy/tuning/json-logs\n")
    (sroot / "opt" / "zeek" / "site").mkdir(parents=True)
    (sroot / "opt" / "zeek" / "site" / "local.zeek").write_text("# comment\n@load base/frameworks\n@load policy/tuning/json-logs\n")
    check("a sensor file equal to deploy/sensors/ (apart from comments) is not drift",
          DD.sensor_drift(gr / "deploy", sroot) == [])
    (sroot / "opt" / "zeek" / "site" / "local.zeek").write_text("@load base/frameworks\n")   # the package's own version
    check("a sensor file replaced by a package update is reported as changed, by its path",
          DD.sensor_drift(gr / "deploy", sroot) == [("changed", "sensors/opt/zeek/site/local.zeek")],
          str(DD.sensor_drift(gr / "deploy", sroot)))
    (sroot / "opt" / "zeek" / "site" / "local.zeek").unlink()
    # soc-nmap: the only thing sudo lets the SOC user run as root for scans (deploy/bin/soc-nmap)
    group("soc-nmap")
    wrap = SUITE / "deploy" / "bin" / "soc-nmap"
    nres = tmp / "nmap_results"
    nres.mkdir()
    (nres / "_targets_20260924.txt").write_text("127.0.0.1\n# a comment\n")   # named like the scans' own lists
    wenv = dict(os.environ, SOC_NMAP_RESULTS=str(nres))
    wrun = lambda *a: subprocess.run([sys.executable, "-I", str(wrap), *a], env=wenv, capture_output=True,  # noqa: E731
                                     text=True, timeout=60)
    (nres / "_vulnbatches_1").mkdir()
    r = wrun("-sn", "-oA", str(nres / "_vulnbatches_1" / "vuln_1"), "-iL", str(nres / "_targets_20260924.txt"))
    check("an allowed scan runs, with its target list and its output inside the results directory",
          r.returncode == 0 and (nres / "_vulnbatches_1" / "vuln_1.xml").exists()
          and "127.0.0.1" in (nres / "_vulnbatches_1" / "vuln_1.xml").read_text(), r.stderr)
    (nres / "shadow.txt").symlink_to("/etc/shadow")
    attacks = [["--script", "/tmp/x.nse", "127.0.0.1"], ["--script-args", "x=1", "127.0.0.1"],
               ["-oA", "/etc/cron.d/x", "127.0.0.1"], ["-iL", "/etc/shadow"], ["-iL", str(nres / "shadow.txt")],
               ["-oA", str(nres / ".." / "x"), "127.0.0.1"], ["-sn", "-oA", str(nres / "a"), "-oA", "/tmp/b", "1.2.3.4"],
               ["--interactive"], ["-sn", ";id"], ["-sn", "--datadir", "/tmp", "127.0.0.1"]]
    refused = [a_ for a_ in attacks if wrun(*a_).returncode != 2]
    check("everything that would make sudo nmap root is refused (scripts by path, script args, writes or reads "
          "outside the results directory, symlinks, repeated -oA, unknown options, shell text)", not refused, str(refused))
    check("the scans call soc-nmap, never nmap itself, and the sudo rule names only soc-nmap",
          not any("sudo nmap " in (SUITE / "kali" / f_).read_text()
                  for f_ in ("1_host_discovery.sh", "2_port_service_scan.sh", "3_vuln_scan.sh"))
          and "/usr/bin/nmap" not in "".join(l_ for l_ in (SUITE / "deploy" / "sudoers" / "soc-nmap-nopasswd").read_text()
                                                  .splitlines() if not l_.startswith("#")))

    group("deploy drift")
    fw_ok = "Status: active\nDefault: deny (incoming)\n\n22/tcp ALLOW IN 10.100.0.0/24 # SSH admin\n"
    (gr / "deploy" / "firewall").mkdir()
    (gr / "deploy" / "firewall" / "ufw-status.txt").write_text("# header\n" + fw_ok)
    check("a live firewall equal to deploy/firewall/ (header and blank lines aside) is not drift",
          DD.firewall_drift(gr / "deploy", live=fw_ok.replace("\n\n", "\n")) == [])
    check("a rule added by hand on the machine is reported as firewall drift",
          DD.firewall_drift(gr / "deploy", live=fw_ok + "3389/tcp ALLOW IN Anywhere\n")
          == [("firewall", "firewall/ufw-status.txt")])
    check("...and the installer's own check ignores the firewall (it is never applied automatically)",
          all(k_ != "firewall" for k_, _ in DD.drift(gr / "deploy", gl, check_crontab=False, sensors_root=None,
                                                    check_firewall=False)))
    (gr / "deploy" / "packages.txt").write_text("# SOC\nnmap   # scans\nzeek\n")
    (gr / "deploy" / "packages-baseline.txt").write_text("# stock\nkali-linux-core\n")
    pdrift = DD.package_drift(gr / "deploy", manual={"nmap", "zeek", "kali-linux-core", "netcat-traditional"},
                              installed={"nmap", "kali-linux-core", "netcat-traditional"})
    check("packages: one listed but not installed is 'package missing', one installed by hand but unlisted is "
          "'package unlisted', stock Kali packages are ignored",
          sorted(pdrift) == [("package missing", "zeek"), ("package unlisted", "netcat-traditional")], str(pdrift))
    check("...and every package in the real deploy/packages.txt is a single package name",
          all(re.fullmatch(r"[a-z0-9][a-z0-9.+-]+", p_) for p_ in DD._pkg_list(SUITE / "deploy" / "packages.txt")))
    check("...and one that is not on the machine at all as not installed",
          DD.sensor_drift(gr / "deploy", sroot) == [("not installed", "sensors/opt/zeek/site/local.zeek")])

    print(json.dumps([{"name": n_, "ok": ok_, "detail": d} for n_, ok_, d in results]))
    return 0


def run_isolated() -> None:
    tmp = tempfile.mkdtemp(prefix="soc_selftest_")
    env = {k: v for k, v in os.environ.items() if k not in ("SOC_INGEST_URL", "NTFY_TOPIC", "SOC_SUPPRESSIONS_FILE")}
    env.update(SOC_DATA_DIR=tmp, SOC_KALI_DIR=tmp, SOC_PLAYBOOKS_DRY_RUN="1")
    shutil_assets = DATA / "assets.json"
    if shutil_assets.exists():
        Path(tmp, "assets.json").write_bytes(shutil_assets.read_bytes())
    try:
        proc = subprocess.run([sys.executable, str(Path(__file__).resolve()), "--inner"], env=env,
                              capture_output=True, text=True, timeout=180)
    except subprocess.TimeoutExpired:
        check("isolated tests finished", False, "timed out after 180 s")
        return
    out = proc.stdout.strip().splitlines()
    try:
        for r in json.loads(out[-1]):
            results.append((r["name"], r["ok"], r["detail"]))
    except (ValueError, IndexError):
        check("isolated tests ran", False, (proc.stderr or proc.stdout)[-400:])
    finally:
        import shutil
        shutil.rmtree(tmp, ignore_errors=True)


# --------------------------------------------------------------------------- #
#  Part 2 -- live, read-only
# --------------------------------------------------------------------------- #

def api_get(path: str, token: str, timeout: int = 20):
    req = urllib.request.Request(API + path, headers={"Authorization": f"Bearer {token}"})
    with urllib.request.urlopen(req, timeout=timeout) as resp:
        return json.load(resp)


def read_token() -> str | None:
    try:
        keys = json.loads((DATA / "api_keys.json").read_text())["keys"]
        return next((k["token"] for k in keys if k["role"] == "read"), keys[0]["token"])
    except (OSError, ValueError, KeyError, IndexError):
        return None


def run_live() -> None:
    sys.path[:0] = [str(SCRIPTS)]
    group("live")
    from service_watchdog import SERVICES
    for svc in SERVICES:
        state = subprocess.run(["systemctl", "is-active", svc], capture_output=True, text=True).stdout.strip()
        check(f"service {svc} is active", state == "active", state)
    from soc_doctor import TIMERS
    bad = [t for t in TIMERS if subprocess.run(["systemctl", "is-enabled", t], capture_output=True, text=True).stdout.strip() != "enabled"]
    check("all scheduled timers are enabled", not bad, ", ".join(bad))
    import deploy_drift
    drifted = deploy_drift.drift()
    for kind, name in drifted:
        check(f"deploy/ matches this machine: {name}", False, f"{kind}: {deploy_drift.EXPLAIN[kind]}")
    check("deploy/ matches the installed units, drop-ins and crontab", not drifted, f"{len(drifted)} difference(s)")

    token = read_token()
    if token is None:
        check("API checks", False, "cannot read data/api_keys.json (run with the soc group: sg soc -c ...)")
    else:
        try:
            health = json.load(urllib.request.urlopen(API + "/api/health", timeout=10))
            check("the API answers /api/health", health.get("status") == "ok")
            for path, key in (("/api/alerts?limit=1", "alerts"), ("/api/incidents", "incidents"), ("/api/entities?limit=1", "entities"),
                              ("/api/metrics", "alerts"), ("/api/sources", "sources"), ("/api/detections", "detectors"),
                              ("/api/suppressions", "rules"), ("/api/playbooks", "playbooks"),
                              ("/api/search?source=alerts&q=selftest&limit=1", "results")):
                try:
                    check(f"API {path.split('?')[0]} responds", key in api_get(path, token))
                except Exception as e:
                    check(f"API {path.split('?')[0]} responds", False, f"{type(e).__name__}: {e}")
        except Exception as e:
            check("the API answers /api/health", False, f"{type(e).__name__}: {e}")

    import source_health
    for s in source_health.check():
        if s["status"] == "stale":
            check(f"data source {s['name']} is fresh", False, f"silent for {s['age_minutes']} min")
    check("no data source is silent", not any(s["status"] == "stale" for s in source_health.check()))

    import correlate, playbooks, suppressions, alert_aging
    for label, errs in (("correlation rules", correlate.load_rules()[2]), ("suppression rules", suppressions.load_rules()[1]),
                        ("playbooks", playbooks.load_playbooks()[1]), ("alert aging rules", alert_aging.load_rules()[1])):
        check(f"{label} file is valid", not errs, "; ".join(errs)[:200])
    try:
        st = json.loads((DATA / "backup_status.json").read_text())
        age_h = (datetime.now(timezone.utc) - datetime.fromisoformat(st["last_ok"])).total_seconds() / 3600
        check("the last backup succeeded within 26 hours", age_h <= 26, f"{age_h:.1f} h ago")
    except (OSError, ValueError, KeyError):
        check("the last backup succeeded within 26 hours", False, "no backup status yet (run kali/backup_data.sh)")
    try:
        meta = json.loads((DATA / "threat_intel" / "meta.json").read_text())
        age_h = (datetime.now(timezone.utc) - datetime.fromisoformat(meta["kev"]["fetched"])).total_seconds() / 3600
        check("the threat-intel feeds were refreshed within 72 hours", age_h <= 72, f"{age_h:.0f} h ago")
    except (OSError, ValueError, KeyError):
        check("the threat-intel feeds were refreshed within 72 hours", False, "never downloaded")
    doc = subprocess.run([sys.executable, str(SCRIPTS / "soc_doctor.py")], capture_output=True, text=True, timeout=120)
    check("soc_doctor reports no failures", doc.returncode == 0, (doc.stdout.strip().splitlines() or [""])[-2][:120])


# --------------------------------------------------------------------------- #
#  Part 3 -- canary through the real pipeline
# --------------------------------------------------------------------------- #

def run_canary() -> None:
    sys.path[:0] = [str(SCRIPTS)]
    group("canary")
    import soc_core
    stamp = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")
    rec = soc_core.emit_alert(soc_core.Alert(
        type="intrusion", severity="normal", detector="selftest", title=f"SELFTEST canary {stamp}",
        description="Harmless round-trip test by soc_selftest.py; resolved automatically.",
        test=True), echo=False)
    check("the canary is marked as a test alert (it stays out of the real feed)", rec.get("test") is True)
    time.sleep(1)
    in_feed = any(a["id"] == rec["id"] for a in soc_core._load_snapshot())
    check("the alert reached the live feed", in_feed)
    token = read_token()
    if token:
        try:
            got = api_get("/api/alerts?limit=20", token)["alerts"]
            check("the alert is served by the API", any(a["id"] == rec["id"] for a in got))
        except Exception as e:
            check("the alert is served by the API", False, str(e))
    try:
        soc_core.set_alert_status(rec["id"], "resolved", note="Self-test canary; resolved automatically.", actor="selftest")
        check("the canary was resolved (dashboard stays clean)", True)
    except ValueError as e:
        check("the canary was resolved (dashboard stays clean)", False, str(e))


def alert_on_change(summary: dict, previous: dict) -> None:
    """One alert when the set of failing checks changes to something non-empty, one 'recovered'
    alert when a failing run is followed by a clean one. The same failures on consecutive
    daily runs stay quiet."""
    sys.path[:0] = [str(SCRIPTS)]
    from soc_core import Alert, emit_alert
    now_failed = sorted(f["name"] for f in summary["failures"])
    was_failed = sorted(f["name"] for f in previous.get("failures", []))
    if now_failed and now_failed != was_failed:
        emit_alert(Alert(
            type="intrusion", severity="medium", detector="selftest",
            title=f"Self-test failed: {len(now_failed)} check(s)",
            description=("The daily self-test found problems in the detection pipeline: "
                         + "; ".join(now_failed[:5]) + (" ..." if len(now_failed) > 5 else "")
                         + ". Run: sg soc -c 'python3 /opt/sentinel-soc/scripts/soc_selftest.py'"),
            details={"failures": summary["failures"][:20], "passed": summary["passed"], "total": summary["total"]}), echo=False)
    elif not now_failed and was_failed:
        emit_alert(Alert(
            type="intrusion", severity="normal", detector="selftest",
            title="Self-test recovered",
            description=f"All {summary['total']} self-test checks pass again.",
            details={"passed": summary["passed"], "total": summary["total"]}), echo=False)


def main() -> int:
    args = sys.argv[1:]
    if "--inner" in args:
        return inner()
    started = time.time()
    if "--live" not in args:
        run_isolated()
    if "--isolated" not in args:
        run_live()
    if "--canary" in args:
        run_canary()
    failed = [r for r in results if not r[1]]
    for name, ok, detail in results:
        print(f"{'PASS' if ok else 'FAIL'}  {name}" + (f"   [{detail}]" if detail and not ok else ""))
    print(f"\n{len(results) - len(failed)}/{len(results)} passed in {time.time() - started:.0f}s"
          + ("" if not failed else f" -- {len(failed)} FAILED"))
    summary = {"at": datetime.now(timezone.utc).isoformat(), "passed": len(results) - len(failed), "total": len(results),
               "failed": len(failed), "failures": [{"name": n, "detail": d} for n, _, d in failed]}
    try:
        previous = json.loads((DATA / "selftest_last.json").read_text())
    except (OSError, ValueError):
        previous = {}
    try:
        DATA.mkdir(parents=True, exist_ok=True)
        (DATA / "selftest_last.json").write_text(json.dumps(summary, indent=2))
    except OSError:
        pass
    if "--alert-on-fail" in args:
        alert_on_change(summary, previous)
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
